"""Phase 7: the right-to-left Word masters.

WHAT THIS PROTECTS

    Word has no CSS, so `dir="rtl"` buys nothing here: a DOCX is right-to-left
    only if the XML says so, in four separate places. Saying only the obvious
    one - `w:bidi` on the section - produces a document that opens, fits, looks
    plausible in a structural test, and lays its Arabic out left-to-right.

    The English masters must not move. They are tuned to fit exactly one page
    (sessions 6 and 8), so the RTL pair is a SEPARATE FILE, not a switch inside
    them, and this file pins that separation.

    `tests/test_docx_validity.py` already covers schema order for all four
    masters, including the elements added here - `w:bidi`, `w:rtl`, `w:bCs`,
    `w:szCs`, `w:bidiVisual` - because it globs the directory and its SEQUENCES
    table knows where each belongs. That is the half that Word, and only Word,
    would otherwise catch.
"""
from __future__ import annotations

import io
import json
import zipfile
from pathlib import Path

import pytest
from docx import Document

from app import registry
from app.exporters.docx import DocxExportError, _master_for, render_docx
from app.labels import _AR, _key

ROOT = Path(__file__).resolve().parent.parent
MASTER_DIR = ROOT / "word_masters"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

EN = json.loads((ROOT / "data" / "sample_resume.json").read_text(encoding="utf-8"))
AR = json.loads((ROOT / "data" / "sample_resume_ar.json").read_text(encoding="utf-8"))

RTL_PAIRS = [("ats_standard.docx", "ats_standard_rtl.docx"),
             ("modern_editorial.docx", "modern_editorial_rtl.docx")]


def _document_xml(blob: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(blob)) as z:
        return z.read("word/document.xml").decode("utf-8")


def _master_bytes(name: str) -> bytes:
    return (MASTER_DIR / name).read_bytes()


# --- master selection -----------------------------------------------------

@pytest.mark.parametrize("key", ["ats-t1", "modern-t1", "ats-t20", "modern-t14"])
def test_an_arabic_resume_gets_the_rtl_master(key):
    tpl = registry.get(key)
    assert _master_for(tpl, EN) == tpl.docx_master
    assert _master_for(tpl, {**EN, "lang": "ar"}).endswith("_rtl.docx")


def test_a_resume_with_no_lang_gets_the_english_master():
    """`lang` is optional and absent-means-English, so every résumé written
    before the bilingual work still exports exactly as it did."""
    tpl = registry.get("ats-t1")
    assert _master_for(tpl, {"name": "A", "title": "B"}) == "ats_standard.docx"


def test_a_missing_rtl_master_raises_instead_of_falling_back(monkeypatch):
    """Falling back to the English master would hand an Arabic résumé a
    left-to-right Word file with English headings - output that is wrong rather
    than absent, which is the failure shape this project keeps finding.
    """
    import app.exporters.docx as mod
    monkeypatch.setattr(mod, "_RTL_MASTERS", {})
    with pytest.raises(DocxExportError) as exc:
        mod._master_for(registry.get("ats-t1"), {**EN, "lang": "ar"})
    assert "right-to-left" in str(exc.value)


# --- the four places RTL has to be said -----------------------------------

@pytest.mark.parametrize("_, rtl", RTL_PAIRS)
def test_the_rtl_master_says_rtl_in_all_four_places(_, rtl):
    xml = _document_xml(_master_bytes(rtl))
    assert "<w:bidi/>" in xml, "no w:bidi at all"
    assert "<w:rtl/>" in xml, (
        "w:bidi without w:rtl: Word lays the glyphs out left-to-right inside a "
        "right-aligned paragraph")
    assert 'w:cs="' in xml, (
        "no complex-script font: Arabic takes its face from w:rFonts/@w:cs, not "
        "@w:ascii, so Word substitutes one")
    assert "<w:bidiVisual/>" in xml, "the chip table does not mirror"


@pytest.mark.parametrize("_, rtl", RTL_PAIRS)
def test_every_paragraph_and_run_is_marked(_, rtl):
    """A post-pass, not a flag threaded through each component, so one
    paragraph cannot be left English-aligned."""
    from lxml import etree
    root = etree.fromstring(_document_xml(_master_bytes(rtl)).encode("utf-8"))
    body = root.find(f"{W}body")
    paras = list(body.iter(f"{W}p"))
    runs = list(body.iter(f"{W}r"))
    assert paras and runs
    unmarked_p = [p for p in paras
                  if p.find(f"{W}pPr") is None
                  or p.find(f"{W}pPr").find(f"{W}bidi") is None]
    unmarked_r = [r for r in runs
                  if r.find(f"{W}rPr") is None
                  or r.find(f"{W}rPr").find(f"{W}rtl") is None]
    assert not unmarked_p, f"{len(unmarked_p)} paragraphs without w:bidi"
    assert not unmarked_r, f"{len(unmarked_r)} runs without w:rtl"


@pytest.mark.parametrize("ltr, _", RTL_PAIRS)
def test_the_english_masters_say_none_of_it(ltr, _):
    """The RTL work must be invisible to English, and the file is the proof."""
    xml = _document_xml(_master_bytes(ltr))
    for tag in ("<w:bidi/>", "<w:rtl/>", "<w:bidiVisual/>"):
        assert tag not in xml, f"{ltr} carries {tag}"


# --- headings and rendering ----------------------------------------------

@pytest.mark.parametrize("key", ["ats-t1", "modern-t1"])
def test_the_rtl_export_carries_arabic_headings(key):
    doc = Document(io.BytesIO(render_docx({**AR, "lang": "ar"}, key)))
    text = "\n".join(p.text for p in doc.paragraphs)
    for label in ("Experience", "Skills", "Education", "Summary", "Profile"):
        assert label not in text, f"{key}: English heading {label!r} in an Arabic export"
    # and at least a few real Arabic headings are present
    wanted = [_AR[_key(x)] for x in ("Experience", "Education", "Skills")]
    assert sum(w in text for w in wanted) >= 2, f"{key}: Arabic headings missing"


@pytest.mark.parametrize("key", ["ats-t1", "modern-t1"])
def test_no_unrendered_tag_survives_in_either_language(key):
    for data in (EN, AR):
        doc = Document(io.BytesIO(render_docx(data, key)))
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "{{" not in text and "{%" not in text


def test_the_arabic_master_avoids_a_font_with_no_arabic():
    """Georgia has no Arabic glyphs at all. Left in place, Word substitutes a
    face of its own choosing - the same silent substitution the HTML side
    solved by self-hosting fonts."""
    xml = _document_xml(_master_bytes("modern_editorial_rtl.docx"))
    assert "Georgia" not in xml
    assert "Georgia" in _document_xml(_master_bytes("modern_editorial.docx"))
