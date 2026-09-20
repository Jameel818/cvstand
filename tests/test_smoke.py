"""Smoke + contract tests. Run: venv/Scripts/python -m pytest -q"""
import re
import dataclasses
import io

import pytest

from app import create_app
from app import registry
from app.rendering import canvas_html
from app.schema import ResumeValidationError, normalize, validate
from tests import samples

SAMPLE = samples.ENGLISH


@pytest.fixture()
def client():
    app = create_app()
    app.config["TESTING"] = True
    return app.test_client()


def test_sample_resume_is_valid():
    validate(SAMPLE)


def test_invalid_resume_reports_all_errors():
    with pytest.raises(ResumeValidationError):
        validate({"title": "no name"})  # missing required 'name'


def test_normalize_fills_optionals_and_caps_chips():
    n = normalize({"name": "X", "title": "Y", "achievements": [
        {"metric": "1", "label": "a"}, {"metric": "", "label": "b"},
        {"metric": "3", "label": "c"}, {"metric": "4", "label": "d"}, {"metric": "5", "label": "e"},
    ]})
    assert [a["metric"] for a in n["achievements"]] == ["1", "3", "4", "5"]  # blank dropped, capped at 4
    assert n["languages"] == [] and n["contact"]["social"] == []


@pytest.mark.parametrize("lang,sample", samples.BOTH)
@pytest.mark.parametrize("key", registry.ported_keys())
def test_ported_templates_render_within_page_box(key, lang, sample):
    """The catalogue-wide render contract, in BOTH directions (phase 8).

    Asserted against the sample's own first skill rather than the literal
    "Brand Systems"/"Expert", so the Arabic pass is a real check and not a
    search for English strings that happen to be absent. The rating shows as
    the level word, or as the percent where the résumé gave one — see
    tests/test_unrated_skill.py for why those are the same claim."""
    html = canvas_html(sample, key)
    assert "width:850px" in html and "height:1100px" in html
    # skill name AND level present as text
    skill = sample["skills"][0]
    assert skill["name"] in html
    assert skill["level"] in html or (
        skill.get("percent") is not None and f"{skill['percent']}%" in html)
    # no CSS pseudo-element content
    assert "::before" not in html and "::after" not in html


@pytest.fixture()
def unported_key(monkeypatch):
    """All 49 catalogued layouts are ported, so synthesise a catalogued-but-
    unported entry to keep that branch covered."""
    key = "ats-t25"
    stub = dataclasses.replace(registry.TEMPLATES[key], ported=False)
    monkeypatch.setitem(registry.TEMPLATES, key, stub)
    return key


def test_unported_template_key_is_rejected(unported_key):
    from app.rendering import UnknownTemplate
    with pytest.raises(UnknownTemplate):
        canvas_html(SAMPLE, unported_key)


def test_pages_ok(client):
    for path in ("/", "/templates", "/templates?cat=modern", "/templates?cat=ats", "/builder"):
        assert client.get(path).status_code == 200


def test_render_and_preview_api(client):
    r = client.post("/api/render", json={"data": SAMPLE, "template_key": "modern-t1"})
    assert r.status_code == 200
    body = r.get_json()
    # Match the class TOKEN, not the whole attribute: 27 of the 49 roots
    # now read `class="tpl cv-sections-taj"` (the Arabic section-face
    # split), and an exact-string check called that a missing canvas.
    assert body["ok"]
    assert re.search(r'<div class="tpl[ "]', body["html"]), body["html"][:200]
    assert body["doc"].startswith("<!DOCTYPE")
    assert client.get("/preview?template_key=ats-t3").status_code == 200


def test_set_template_rejects_unported(client, unported_key):
    assert client.post("/api/template", json={"template_key": unported_key}).status_code == 409
    assert client.post("/api/template", json={"template_key": "nope"}).status_code == 400


DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


@pytest.fixture()
def stored_resume():
    """Put a résumé in the store for one test, then put the store back.

    `/export/*` reads the saved résumé — there is no way to hand it data — so
    a route-level test of the Arabic path has to write one. The whole session
    shares ONE data dir (tests/conftest.py), so without the restore this test
    leaves an Arabic résumé behind and the next test that reads the store gets
    a document it never asked for. That is not hypothetical: it is what this
    fixture was written for, after
    `test_ui_language.py::test_interface_language_does_not_touch_the_document`
    started failing in the full run while passing alone."""
    from app.config import RESUME_PATH

    before = RESUME_PATH.read_bytes() if RESUME_PATH.exists() else None

    def _put(client, data):
        assert client.put("/api/resume", json=data).status_code == 200

    yield _put

    if before is None:
        RESUME_PATH.unlink(missing_ok=True)
    else:
        RESUME_PATH.write_bytes(before)


@pytest.mark.parametrize("lang,sample", samples.BOTH)
@pytest.mark.parametrize("key", ["modern-t1", "ats-t1"])
def test_docx_export_returns_a_word_file(client, stored_resume, key, lang, sample):
    """Both directions through the route, not just the exporter: an Arabic
    résumé takes a different master, and a missing one raises rather than
    falling back (phase 7), so the route is where that becomes a 501."""
    stored_resume(client, sample)
    resp = client.get(f"/export/docx?template_key={key}")
    assert resp.status_code == 200
    assert resp.mimetype == DOCX_MIME
    assert resp.data[:2] == b"PK"  # a .docx is a zip package


def _docx(key, data=None):
    from docx import Document
    from app.exporters.docx import render_docx
    return Document(io.BytesIO(render_docx(data if data is not None else SAMPLE, key)))


@pytest.mark.parametrize("key,profile_head", [("ats-t1", "Summary"), ("modern-t1", "Profile")])
def test_docx_master_fills_every_section(key, profile_head):
    paras = [p.text for p in _docx(key).paragraphs]
    body = "\n".join(paras)
    # no unrendered docxtpl tags survived
    assert "{{" not in body and "{%" not in body
    for head in (profile_head, "Key Achievements", "Experience", "Skills",
                 "Education", "Certifications", "Tools", "Languages"):
        assert head in paras, f"{key}: missing section {head!r}"
    # every experience entry and bullet made it through the loops
    assert sum(1 for p in paras if p.startswith("Halden & Row")) == 1, "'&' must survive"
    assert len([p for p in _docx(key).paragraphs if p.style.name == "List Bullet"]) == 6
    # skill level word AND dot glyphs are both real text
    assert any("Brand Systems — Expert" in p and "●" in p for p in paras)


def test_docx_unmapped_level_gets_no_dots():
    """"Native" is off the LEVEL_DOTS scale — it must render as text only, never
    as five empty circles contradicting the word beside it."""
    langs = [p.text for p in _docx("ats-t1").paragraphs if p.text.startswith(("English", "German"))]
    assert langs and all("○" not in t and "●" not in t for t in langs)


@pytest.mark.parametrize("key", ["modern-t1", "ats-t1"])
def test_docx_empty_sections_leave_no_orphan_headings(key):
    minimal = {"name": "Dana Ortiz", "contact": {"email": "d@example.com"}}
    doc = _docx(key, minimal)
    paras = [p.text for p in doc.paragraphs if p.text.strip()]
    for head in ("Summary", "Profile", "Key Achievements", "Experience",
                 "Skills", "Education", "Certifications", "Tools", "Languages"):
        assert head not in paras, f"{key}: orphan heading {head!r} with no content"
    assert not doc.tables, "chip table must go when there are no achievements"


def _docx_images(key, data):
    return [r for r in _docx(key, data).part.rels.values() if "image" in r.reltype]


@pytest.fixture()
def uploaded_photo(client):
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (600, 600), (180, 60, 40)).save(buf, "PNG")
    buf.seek(0)
    resp = client.post("/api/photo", data={"photo": (buf, "me.png")},
                       content_type="multipart/form-data")
    return resp.get_json()["url"]


def test_docx_modern_embeds_the_photo(uploaded_photo):
    assert len(_docx_images("modern-t1", dict(SAMPLE, photo_url=uploaded_photo))) == 1


def test_docx_ats_never_embeds_a_photo(uploaded_photo):
    """No ATS template has a photo slot — images defeat resume parsers — so the
    ATS master must not grow one even when a photo is set."""
    assert _docx_images("ats-t1", dict(SAMPLE, photo_url=uploaded_photo)) == []


@pytest.mark.parametrize("url", ["", "/uploads/does_not_exist.jpg"])
def test_docx_photo_absent_or_dangling_leaves_no_image(url):
    """A deleted upload must degrade to no photo, not a broken image or a gap."""
    assert _docx_images("modern-t1", dict(SAMPLE, photo_url=url)) == []


def test_docx_chip_hides_metric_and_caption_together():
    """A chip past the filled count must lose its caption too — never a caption
    floating with no number."""
    data = {"name": "Dana Ortiz", "contact": {"email": "d@example.com"},
            "achievements": [{"metric": "9", "label": "Teams"}]}
    cells = [[q.text for q in c.paragraphs] for c in _docx("ats-t1", data).tables[0].row_cells(0)]
    assert cells[0] == ["9", "Teams"]
    assert all(all(t == "" for t in cell) for cell in cells[1:])
