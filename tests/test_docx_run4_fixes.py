"""Run 4 item 6 fixes that hold for every Word design:

* a vertically merged cell's height is shared by the rows it spans - the
  estimator read modern-t14's header (photo merged over three rows) as 415pt
  instead of its 249pt, and the near-one-page fit then squeezed the column;
* in a right-to-left (bidiVisual) table Word draws a cell's w:left / w:right
  BORDER at its leading / trailing edge, so a logical "start" border is
  w:left there too (the Arabic timeline rails sat away from their nodes)."""
from __future__ import annotations

import io
import json
import zipfile

from lxml import etree

from app.exporters.docx import render_docx

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
EN = json.loads(open("data/demo_resume.json", encoding="utf-8").read())
AR = json.loads(open("data/demo_resume_ar.json", encoding="utf-8").read())


def _doc(data, key):
    return etree.fromstring(zipfile.ZipFile(io.BytesIO(render_docx(data, key)))
                            .read("word/document.xml"))


def test_a_merged_cell_spans_its_rows():
    from docx import Document
    from app.exporters.docx_measure import Measure
    d = Document()
    t = d.add_table(rows=3, cols=2)
    a, b = t.rows[0].cells[1], t.rows[2].cells[1]
    a.merge(b)
    for _ in range(9):
        a.add_paragraph("x")
    t.rows[1].cells[0].paragraphs[0].add_run("y")
    from app.exporters.docx_theme import resolve
    m = Measure(resolve(EN, "modern-t14"))
    whole = m.table(t._tbl)
    merged = m.block(a._tc, 100)
    assert abs(whole - merged) < 1.0, (whole, merged)


def test_t14_rails_are_on_the_node_side_in_both_languages():
    for data, edge in ((EN, "left"), (AR, "left")):
        root = _doc(data, "modern-t14")
        sides = {etree.QName(e).localname for tc in root.iter(W + "tc")
                 for b in [tc.find(W + "tcPr/" + W + "tcBorders")] if b is not None
                 for e in b if e.get(W + "color") == "8A8A93"}
        assert sides == {edge}, sides
