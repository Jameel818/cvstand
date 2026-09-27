"""Word files carry the fonts they use (typography step 6, spec §6.3 / §7.7).

Read straight from the package: the settings switch, the font table, the
relationships, the content type, the obfuscated parts. The round trip -
de-obfuscating each part with its own w:fontKey gives back the source TTF
byte for byte - is what proves the key derivation. That Word itself then
DRAWS with these fonts is proved in real Word by
tools/verify_word_embedding.py (it cannot run in CI).
"""
from __future__ import annotations

import io
import zipfile

import docx
import pytest
from fontTools.ttLib import TTFont
from lxml import etree

from app.exporters import docx_font_embed as E
from app.exporters import docx_theme as T
from app.exporters.docx import render_docx
from app.typography import face
from app.typography.faces import FONT_DIR
from tests.samples import ARABIC, ENGLISH

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
PR = "{http://schemas.openxmlformats.org/package/2006/relationships}"

#: One EN and one AR default, a non-RIBBI 800, Almarai's Bold-named family,
#: the Reserved-Font-Name families shipped unmodified (IBM Plex Sans Arabic,
#: Source Sans 3), a template default built in step 6 (Open Sans) and a
#: chosen set in each language.
CASES = [
    ("en-default-t16", ENGLISH, "modern-t16"),          # Montserrat 800 + Open Sans
    ("ar-default-t3", ARABIC, "modern-t3"),              # Tajawal + IBM Plex Sans Arabic
    ("en-source-sans", ENGLISH, "ats-t9"),               # Source Sans 3 (unmodified)
    ("en-chosen-800", dict(ENGLISH, font_heading="Montserrat", font_heading_weight=800,
                           font_body="Inter", font_body_weight=300), "ats-t1"),
    ("ar-almarai", dict(ARABIC, font_name="Almarai", font_name_weight=700,
                        font_heading="Cairo", font_heading_weight=900,
                        font_body="Amiri"), "modern-t1"),
    ("ar-plex-light", dict(ARABIC, font_body="IBM Plex Sans Arabic",
                           font_body_weight=300), "ats-t1"),
]
IDS = [c[0] for c in CASES]


def _pkg(blob: bytes) -> dict[str, bytes]:
    with zipfile.ZipFile(io.BytesIO(blob)) as z:
        assert z.testzip() is None
        return {n: z.read(n) for n in z.namelist()}


def _drawn(data, key) -> set[tuple[str, int]]:
    """What the document's text draws, recomputed from the RENDERED file."""
    d = docx.Document(io.BytesIO(render_docx(data, key)))
    return T.faces_drawn(d, T.resolve(data, key))


def _embedded(parts) -> list[dict]:
    """[{name, slot, key, part, bytes}] for every embedded font."""
    table = etree.fromstring(parts["word/fontTable.xml"])
    rels = {r.get("Id"): r.get("Target")
            for r in etree.fromstring(parts["word/_rels/fontTable.xml.rels"])}
    out = []
    for font in table.iter(f"{W}font"):
        for tag, slot in (("embedRegular", "regular"), ("embedBold", "bold")):
            el = font.find(f"{W}{tag}")
            if el is None:
                continue
            part = "word/" + rels[el.get(f"{R}id")]
            out.append({"name": font.get(f"{W}name"), "slot": slot,
                        "key": el.get(f"{W}fontKey"), "part": part, "bytes": parts[part]})
    return out


# ---- the key derivation ---------------------------------------------------------

def test_the_key_is_the_guid_bytes_reversed():
    guid = "{00112233-4455-6677-8899-AABBCCDDEEFF}"
    assert E.font_key(guid) == bytes.fromhex("FFEEDDCCBBAA99887766554433221100")


def test_only_the_first_32_bytes_change():
    data = bytes(range(64)) * 4
    guid = E.new_guid()
    ob = E.obfuscate(data, guid)
    assert ob[32:] == data[32:] and ob[:32] != data[:32]
    assert E.obfuscate(ob, guid) == data


# ---- the package ----------------------------------------------------------------

@pytest.mark.parametrize("_id,data,key", CASES, ids=IDS)
def test_settings_ask_word_to_use_embedded_fonts(_id, data, key):
    root = etree.fromstring(_pkg(render_docx(data, key))["word/settings.xml"])
    kids = [etree.QName(c).localname for c in root if isinstance(c.tag, str)]
    assert "embedTrueTypeFonts" in kids
    i = kids.index("embedTrueTypeFonts")
    assert set(kids[:i]) <= set(E.SETTINGS_BEFORE), f"out of schema order: {kids[:i + 1]}"
    assert "saveSubsetFonts" not in kids, "a subset would stop the recipient editing"


@pytest.mark.parametrize("_id,data,key", CASES, ids=IDS)
def test_exactly_the_drawn_faces_are_embedded_and_wired(_id, data, key):
    parts = _pkg(render_docx(data, key))
    emb = _embedded(parts)
    want = {(face(*f)["word_family_name"], "bold" if face(*f)["word_bold"] else "regular")
            for f in _drawn(data, key)}
    assert {(e["name"], e["slot"]) for e in emb} == want, "missing or extra faces"
    odttf = {n for n in parts if n.endswith(".odttf")}
    assert odttf == {e["part"] for e in emb}, "an orphan or missing font part"
    ct = parts["[Content_Types].xml"].decode()
    assert 'Extension="odttf"' in ct and E.ODTTF_TYPE in ct
    assert "/word/fontTable.xml" in ct


@pytest.mark.parametrize("_id,data,key", CASES, ids=IDS)
def test_round_trip_gives_back_the_source_ttf(_id, data, key):
    parts = _pkg(render_docx(data, key))
    for e in _embedded(parts):
        plain = E.obfuscate(e["bytes"], e["key"])
        sources = [f for f in (face(fam, w) for fam, w in _drawn(data, key))
                   if f["word_family_name"] == e["name"]
                   and bool(f["word_bold"]) == (e["slot"] == "bold")]
        assert len(sources) == 1
        assert plain == (FONT_DIR / sources[0]["ttf"]).read_bytes(), e["name"]


@pytest.mark.parametrize("_id,data,key", CASES, ids=IDS)
def test_embedded_names_match_the_file_and_the_runs(_id, data, key):
    """nameID 1 of the embedded font == its <w:font w:name> == the rFonts the
    styles ask for. Any mismatch and Word ignores the embedded font."""
    parts = _pkg(render_docx(data, key))
    styles = etree.fromstring(parts["word/styles.xml"])
    asked = {f.get(f"{W}ascii") for f in styles.iter(f"{W}rFonts")}
    for e in _embedded(parts):
        font = TTFont(io.BytesIO(E.obfuscate(e["bytes"], e["key"])))
        # The WINDOWS record (3, 1, en-US) - what Word reads. Some official
        # statics also carry a Mac record naming the bare family ("Tajawal"
        # beside "Tajawal ExtraBold"); getDebugName() may return that one.
        rec = font["name"].getName(1, 3, 1, 0x409)
        assert rec is not None and rec.toUnicode() == e["name"]
        assert e["name"] in asked, f"{e['name']} embedded but no style asks for it"


@pytest.mark.parametrize("_id,data,key", CASES, ids=IDS)
def test_font_table_follows_schema_order_and_scripts(_id, data, key):
    parts = _pkg(render_docx(data, key))
    table = etree.fromstring(parts["word/fontTable.xml"])
    family_of = {face(*f)["word_family_name"]: f[0] for f in _drawn(data, key)}
    checked = 0
    for font in table.iter(f"{W}font"):
        name = font.get(f"{W}name")
        if name not in family_of:
            continue
        kids = [etree.QName(c).localname for c in font]
        pos = [E.FONT_ORDER.index(k) for k in kids]
        assert pos == sorted(pos), kids
        # Arabic families are ARABIC_CHARSET (B2), Latin ones ANSI (00).
        want = "B2" if E._is_arabic(family_of[name]) else "00"
        assert font.find(f"{W}charset").get(f"{W}val") == want, name
        checked += 1
    assert checked == len(family_of)


@pytest.mark.parametrize("_id,data,key", CASES, ids=IDS)
def test_python_docx_still_opens_it(_id, data, key):
    d = docx.Document(io.BytesIO(render_docx(data, key)))
    assert any(p.text.strip() for p in d.paragraphs)


# ---- refusals ---------------------------------------------------------------------

def test_a_font_that_forbids_embedding_is_refused(monkeypatch):
    real = E._os2

    def restricted(rel):
        return dict(real(rel), fs_type=2)
    monkeypatch.setattr(E, "_os2", restricted)
    with pytest.raises(E.FontEmbedError, match="fsType 2"):
        E.plan({("Inter", 400)})


def test_two_faces_under_one_word_name_are_refused():
    """Archivo 900 and the Archivo Black family are both "Archivo Black"."""
    assert face("Archivo", 900)["word_family_name"] == face("Archivo Black", 400)["word_family_name"]
    with pytest.raises(E.FontEmbedError, match="two faces"):
        E.plan({("Archivo", 900), ("Archivo Black", 400)})


def test_a_modern_file_does_not_embed_an_unused_role():
    """Modern never uses CV Body Bold; its face must not ride along."""
    data = dict(ENGLISH, font_body="Inter", font_body_weight=300)
    drawn = _drawn(data, "modern-t1")
    assert ("Inter", 700) not in drawn and ("Inter", 300) in drawn


def test_no_file_is_unreasonably_large():
    """Full script-subset TTFs, never trimmed (§6.3) - so bounded, not tiny.
    The heaviest real case (two Naskh faces) stays well under 2.5 MB."""
    data = dict(ARABIC, font_name="Scheherazade New", font_heading="Scheherazade New",
                font_body="Scheherazade New")
    assert len(render_docx(data, "modern-t1")) < 2_500_000
