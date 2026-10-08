"""modern-t18 — Navy frame & brass: a navy band across the main column's top
(312px: caps name, tracked brass title, a translucent rule, the summary) and a
navy rail under a white top-left corner holding a white-framed photo that
straddles the rail's start; the rail lists labelled contact fields and
education (date above degree); the white main column opens with brass Anton
stats over a rule, then dated experience (date above a caps title) and Skills
pinned to the column foot as name | rounded bar | value. From modern/t18.j2.

Simplified in Word: the photo frame (34-334px) sits in a white block, so the
rail's navy begins under the frame (334px) instead of at 312px beside it."""
from __future__ import annotations

from docx.shared import Pt

from ..docx_design import (
    Box, Ctx, bar_shape, design, fmt_cell, fmt_p, page_rects, photo_run, pt, rated, row_height,
    run, skill_pct, tiny, tw,
)
from .common import SidebarPage, Stack, block, bullets, date_range, joined, text_px

NAVY, BRASS, KICK, RAIL_RULE, BAND_RULE = "002B66", "C08A2E", "C9973A", "617CA0", "6B84A6"
VAL, INK, BODY, RULE, TRACK = "E3E9F2", "26303C", "39424E", "C9C6BC", "DCDEE4"
SIDE_PX, BAND = 372, 312


@design("modern-t18")
def build(ctx: Ctx) -> None:
    r, t = ctx.r, ctx.t
    sec = ctx.doc.sections[0]
    sec.top_margin = Pt(0.1)
    sec.bottom_margin = Pt(pt(24))
    page = SidebarPage(ctx, SIDE_PX, side_fill=NAVY, side_pad=(0, 0, 24, 0),
                       main_pad_x=(40, 40), full_height_fill=False)
    page_rects(ctx, [(0, SIDE_PX + 1, 0, 1100, NAVY)])
    side = page.side
    H = "role" if ctx.rtl else "bold"

    # the white corner with the framed photo (white 12px frame at 30px in)
    wb = block(ctx, side, SIDE_PX, fill="FFFFFF", pad=(34, 30, 0, 30))
    photo_run(ctx, wb.p(line=1.0), size_px=312, height_px=300, shape="rect", ring_px=12,
              ring="FFFFFF", placeholder="C8CEDA")
    wb.finish()

    def lhead(b, label, first):
        run(ctx, b.p(before=38 if first else 26, ind_start=42, ind_end=42), label, "heading",
            size=23, color="FFFFFF", spacing=-0.2)
        rule = b.p()
        tiny(rule)
        fmt_p(ctx, rule, before=11, after=15, ind_start=42, ind_end=42,
              border={"bottom": (1, RAIL_RULE, 0)})

    first = [True]

    def section(label, items, write):
        if not items:
            return
        st = Stack(ctx, side, SIDE_PX)
        for i, it in enumerate(items):
            b = st.row()
            if i == 0:
                lhead(b, label, first[0])
                first[0] = False
            write(b, i, it)
        st.done()

    def field(b, i, item):
        label, value = item
        run(ctx, b.p(before=12 if i else 0, ind_start=42, ind_end=42), label, "bold", size=9.5,
            color=KICK, spacing=1.6, caps=True)
        run(ctx, b.p(before=3, ind_start=42, ind_end=42), value, size=12.5, color=VAL)

    def education(b, i, ed):
        dates = date_range(ed["start"], ed["end"])
        if dates:
            run(ctx, b.p(before=18 if i else 0, ind_start=42), dates, size=11, color="9FB0C8")
        run(ctx, b.p(before=2 if dates else (18 if i else 0), ind_start=42, ind_end=42),
            ed["degree"], H, size=14, color="FFFFFF")
        if ed["school"]:
            run(ctx, b.p(ind_start=42, ind_end=42), ed["school"], size=12.5, color="C4D0E2")
        if ed["bullets"]:
            b.pad_top = 8
            bullets(ctx, b, ed["bullets"], size=12, color="DBE3EE", line=1.55, indent=15,
                    ind_start=42)

    c = r["contact"]
    items = [(t("Email"), c["email"]), (t("Phone"), c["phone"]), (t("Location"), c["address"]),
             (t("Portfolio"), c["site"])]
    items += [(t("Link"), s.get("label")) for s in c["social"] if s.get("label")]
    section(t("Contact"), [x for x in items if x[1]], field)
    section(t("Education"), r["education"], education)

    # ---- main column: the navy band (a row at least 312px tall)
    T = page.main_text_px
    band = page.main_row(pad_top=38, fill=NAVY)
    row_height(page.tbl.rows[0], BAND)
    run(ctx, band.p(line=1.04), r["name"], "name", size=44, color="FFFFFF", caps=True,
        spacing=-1)
    if r["title"]:
        run(ctx, band.p(before=10), r["title"], "bold", size=14, color=BRASS, caps=True,
            spacing=3)
    if r["summary"]:
        rule = band.p()
        tiny(rule)
        fmt_p(ctx, rule, before=16, after=13, border={"bottom": (1, BAND_RULE, 0)})
        run(ctx, band.p(line=1.65), r["summary"], size=12, color="D2DCEA")

    if r["achievements"]:
        box = page.main_row(pad_top=24)
        ach = [a for a in r["achievements"] if a["metric"]]
        n = max(1, len(ach))
        tbl = box.table([tw(T) // n] * n, borders={"bottom": (1, "DDDBD3")})
        for k, (cell, a) in enumerate(zip(tbl.rows[0].cells, ach)):
            fmt_cell(ctx, cell, pad=(0, 0, 0, 14 if k < n - 1 else 0))
            cb = Box(ctx, cell, 0, pad_bottom=12)
            # number and label centred in the chip, as in the PDF (user, 2026-10-08)
            run(ctx, cb.p(align="center", line=1.0), a["metric"], "metric", size=24, color=BRASS)
            run(ctx, cb.p(align="center", before=5), a["label"], "bold", size=8.5, color="6E7787", spacing=1,
                caps=True)
            cb.finish()

    def rhead(box, label):
        run(ctx, box.p(), label, "heading", size=23, color=NAVY, spacing=-0.2)
        rule = box.p()
        tiny(rule)
        fmt_p(ctx, rule, before=10, after=14, border={"bottom": (1, RULE, 0)})

    for j, job in enumerate(r["experience"]):
        box = page.main_row(pad_top=16 if (j == 0 and r["achievements"]) else
                            (24 if j == 0 else 13))
        if j == 0:
            rhead(box, t("Work Experience"))
        dates = date_range(job["start"], job["end"])
        if dates:
            run(ctx, box.p(), dates, size=11, color="8B93A0")
        run(ctx, box.p(before=2 if dates else 0), job["role"], H, size=14, color=NAVY,
            caps=True, spacing=0.3)
        meta = joined([job["company"], job["location"]], " · ")
        if meta:
            run(ctx, box.p(), meta, size=12.5, color="4C5666")
        bullets(ctx, box, job["bullets"], size=12, color=BODY, line=1.55, before=7)

    if r["skills"]:
        box = page.main_row(pad_top=16)
        page.push_to_bottom(box, bottom_px=0)      # margin-top:auto in the PDF
        rhead(box, t("Skills"))
        # the value column is 64px; the longest level word widens it (Word
        # would break "Foundational" mid-word in a fixed cell)
        lw = max([64] + [text_px(ctx, "bold", skill_pct(sk)[1], 10.5, 0.3) + 2
                         for sk in r["skills"] if rated(sk)])
        bw = T - 124 - lw - 24
        # ONE table, a row per skill (back-to-back tables would merge in Word)
        tbl = box.table([tw(124 + 12), tw(bw), tw(12 + lw)], rows=len(r["skills"]))
        for i, sk in enumerate(r["skills"]):
            a, b, c3 = tbl.rows[i].cells
            for cell in (a, b, c3):
                fmt_cell(ctx, cell, valign="center", pad=(8 if i else 0, 0, 0, 0))
            run(ctx, fmt_p(ctx, a.paragraphs[0]), sk["name"], "bold", size=12.5, color=NAVY)
            if rated(sk):
                pct, label = skill_pct(sk)
                bar_shape(ctx, fmt_p(ctx, b.paragraphs[0], line=1.0), pct, bw, on=NAVY,
                          off=TRACK, height=9, radius=4.5, fill_round=True)
                run(ctx, fmt_p(ctx, c3.paragraphs[0], align="end"), label, "bold", size=10.5,
                    color=NAVY, spacing=0.3)
    page.finish()
