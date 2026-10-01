"""modern-t6 — Open, centred, brown accent: a centred tracked caps name, the
title in tracked brown caps between two black rules, a centred contact line,
centred brown stats, the summary centred under a hairline; then two equal
columns split by a hairline at the second's start - Experience (brown meta
line), Recognition as a bulleted list, Systems in two columns | Skills as
stacked brown bars with the value as text, Education, Languages. Section
heads are tracked caps ruled underneath. From modern/t6.j2."""
from __future__ import annotations

from docx.shared import Pt

from ..docx_design import (
    Box, Ctx, bar, design, fmt_cell, fmt_p, fmt_table, pt, rated, run, skill_pct, tiny, tw,
)
from .common import SidebarPage, Stack, bullets, date_range, joined

INK, BROWN, GREY, BODY, RULE, TRACK = "1A1A1A", "9C5A2A", "5A5A5A", "333333", "D5D2CE", "E6E3DF"
PAD = 56
COL = (850 - 2 * PAD - 34 - 35) / 2            # each column's text width (334.5px)


@design("modern-t6")
def build(ctx: Ctx) -> None:
    r, t = ctx.r, ctx.t
    sec = ctx.doc.sections[0]
    sec.top_margin = Pt(pt(PAD))
    sec.bottom_margin = Pt(pt(48))
    doc = ctx.doc
    H = "role" if ctx.rtl else "heading"

    def body_p(**kw):
        kw.setdefault("ind_start", PAD)
        kw.setdefault("ind_end", PAD)
        return fmt_p(ctx, doc.add_paragraph(), **kw)

    run(ctx, body_p(align="center", line=1.05), r["name"], "name", size=46, color=INK, caps=True,
        spacing=5)
    if r["title"]:
        # the 9px padding is the borders' own space (inside the paragraph)
        p = body_p(align="center", before=14,
                   border={"top": (1, INK, pt(9)), "bottom": (1, INK, pt(9))})
        run(ctx, p, r["title"], size=14, color=BROWN, caps=True, spacing=5.4)
    c = r["contact"]
    bits = [x for x in (c["email"], c["phone"], c["address"], c["site"]) if x]
    bits += [s.get("label") for s in c["social"] if s.get("label")]
    if bits:
        run(ctx, body_p(align="center", before=12), " · ".join(bits),
            size=12, color="4A4A4A", spacing=0.4)
    if r["achievements"]:
        gap = doc.add_paragraph()
        tiny(gap, before_px=18)
        n = len(r["achievements"])
        Wd = 850 - 2 * PAD
        tbl = doc.add_table(rows=1, cols=n)
        fmt_table(tbl, [tw(Wd) // n] * n, ind=tw(PAD))
        for k, (cell, a) in enumerate(zip(tbl.rows[0].cells, r["achievements"])):
            fmt_cell(ctx, cell, pad=(0, 9, 0, 9))
            cb = Box(ctx, cell, 0)
            run(ctx, cb.p(align="center", line=1.0), a["metric"], "metric", size=28, color=BROWN)
            run(ctx, cb.p(align="center", before=6), a["label"], "bold", size=9, color=GREY,
                spacing=1.3, caps=True)
            cb.finish()
        tail = doc.add_paragraph()
        tiny(tail)
    if r["summary"]:
        p = body_p(align="center", before=18, line=1.7, border={"top": (1, RULE, pt(18))})  # gap 18
        run(ctx, p, r["summary"], size=12.5, color="2E2E2E")

    # ---- the two columns: [56 | 334.5 | 34] [rule 1 | 34 | 334.5 | 56]
    page = SidebarPage(ctx, 850 - (PAD + COL + 34), side_fill=None,
                       side_pad=(0, 35, 0, PAD), main_pad_x=(PAD, 34),
                       full_height_fill=False, side_end=True)
    fmt_cell(ctx, page.tbl.rows[0].cells[1], borders={"start": (1, RULE)})
    side = page.side

    def sh(box, label, before, role="heading"):
        p = box.p(before=before, after=13, border={"bottom": (1, INK, pt(8))})
        run(ctx, p, label, role, size=14, color=INK, caps=True, spacing=2.6)
        return p

    heads = []
    first = [True]

    def section(label, items, write):
        if not items:
            return
        st = Stack(ctx, side, COL)
        for i, it in enumerate(items):
            b = st.row()
            if i == 0:
                hp = sh(b, label, 18 if first[0] else 18)
                if not first[0]:
                    heads.append(hp)
                first[0] = False
            write(b, i, it)
        st.done()

    def skill(b, i, sk):
        p = b.p(before=11 if i else 0, tabs=[(COL, "end", None)])
        run(ctx, p, sk["name"], "bold", size=12.5, color=INK)
        if rated(sk):
            pct, label = skill_pct(sk)
            run(ctx, p, "\t", size=10.5)
            run(ctx, p, label, size=10.5, color=GREY)
            bar(ctx, b, pct, COL, on=BROWN, off=TRACK, height=8, before=5)

    def education(b, i, ed):
        run(ctx, b.p(before=10 if i else 0, line=1.5), ed["degree"], "bold", size=12.5,
            color=INK)
        sub = joined([ed["school"], date_range(ed["start"], ed["end"])], " · ")
        if sub:
            run(ctx, b.p(line=1.5), sub, size=12.5, color=GREY)
        if ed["gpa"]:
            run(ctx, b.p(line=1.5), f"{t('GPA')} {ed['gpa']}", size=12.5, color=GREY)

    section(t("Skills"), r["skills"], skill)
    section(t("Education"), r["education"], education)
    section(t("Languages"), [joined([x["name"], x["level"]], " — ") for x in r["languages"]],
            lambda b, i, s: run(ctx, b.p(line=1.7), s, size=12.5, color=BODY))
    page.fit_side(heads)

    # ---- the first column
    gap = [18]
    for j, job in enumerate(r["experience"]):
        box = page.main_row(pad_top=gap[0] if j == 0 else 16)
        if j == 0:
            sh(box, t("Experience"), 0)
        run(ctx, box.p(), job["role"], H, size=14.5, color=INK)
        meta = joined([job["company"], job["location"],
                       date_range(job["start"], job["end"], " — ")], " · ")
        if meta:
            run(ctx, box.p(before=2, after=2), meta, size=12, color=BROWN)
        bullets(ctx, box, job["bullets"], size=12.5, color=BODY, line=1.6, before=6)
    gap[0] = 22 if r["experience"] else 18
    if r["recognition"]:
        box = page.main_row(pad_top=gap[0])
        sh(box, t("Recognition"), 0)
        bullets(ctx, box, [joined([x["title"], x["detail"]], " — ") for x in r["recognition"]],
                size=12.5, color=BODY, line=1.6)
        gap[0] = 22
    if r["tools"]:
        box = page.main_row(pad_top=gap[0])
        # "Systems" is not a .cv-section in the PDF (Arabic: the body face)
        sh(box, t("Systems"), 0, role="heading" if not ctx.rtl else "role")
        tools = r["tools"]
        col = (COL - 18) / 2
        tbl = box.table([tw(col + 18), tw(col)], rows=(len(tools) + 1) // 2)
        for i, tool in enumerate(tools):
            cell = tbl.rows[i // 2].cells[i % 2]
            fmt_cell(ctx, cell, pad=(0, 0, 0, 18 if i % 2 == 0 else 0))
            cb = Box(ctx, cell, 0, pad_top=7 if i >= 2 else 0)
            run(ctx, cb.p(), tool, size=12.5, color=BODY)
            cb.finish()
        if len(tools) % 2:
            Box(ctx, tbl.rows[-1].cells[1], 0).finish()
    page.finish()
