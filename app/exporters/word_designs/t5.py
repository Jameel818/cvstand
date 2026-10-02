"""modern-t5 — Rounded Dark: a #F2F2F2 page frame (24px) holding an inset navy
sidebar - round photo, white rounded pills (Contact, Summary), dark pills
(Education, Software, Languages), grey education cards with a stacked date -
and a main column with a two-line caps name, coral stats between rules, dark
rounded section pills, a dated timeline (date column, black rail, black node),
a 3-column ruled grid of coral sliders with the percent as text, and
Certifications. From modern/t5.j2."""
from __future__ import annotations

from docx.shared import Pt

from ..docx_design import (
    Box, Ctx, design, fmt_cell, fmt_p, page_rects, photo_run, pt, rated, row_height, run,
    skill_pct, slider_shape, tw, vml_oval,
)
from .common import SidebarPage, Stack, bullets, joined, pill, rounded_block, text_px

FRAME, NAVY, CORAL, TRACK, INK = "F2F2F2", "003366", "F26A5F", "F4D0CC", "1A1A1A"
CARD, MUTED, RULE = "E6E8EC", "555555", "D9D9D9"
PAD, SIDE_W, GAP = 24, 280, 26          # the frame, the navy column, the gutter


@design("modern-t5")
def build(ctx: Ctx) -> None:
    r, t = ctx.r, ctx.t
    sec = ctx.doc.sections[0]
    sec.top_margin = Pt(pt(PAD))
    sec.bottom_margin = Pt(pt(PAD + 32))
    page = SidebarPage(ctx, SIDE_W, side_fill=NAVY, side_pad=(32, 26, 32, 26),
                       main_pad_x=(GAP, PAD), full_height_fill=False, ind_px=PAD)
    # the grey frame and the inset navy column, on every page
    page_rects(ctx, [(0, 850, 0, 1100, FRAME), (PAD, SIDE_W, PAD, 1100 - 2 * PAD, NAVY)])
    side = page.side
    W = SIDE_W - 52
    heads = []
    photo_run(ctx, side.p(align="center"), size_px=150, placeholder="D6DDE3")

    def lines_section(label, lines, *, fill, color, size=16, cv_section=True, line=1.7,
                      text_size=12):
        if not lines:
            return
        st = Stack(ctx, side, W)
        role = "heading" if cv_section or not ctx.rtl else "role"
        for i, s in enumerate(lines):
            b = st.row()
            if i == 0:
                heads.append(pill(ctx, b, label, W, fill=fill, color=color, size=size, before=20,
                                  after=20, spacing=0.5 if fill == INK else 0, role=role,
                                  bg=NAVY).spacer)
            run(ctx, b.p(line=line, ind_start=6), s, size=text_size, color="FFFFFF")
        st.done()

    c = r["contact"]
    contact = [x for x in (c["phone"], c["email"], c["address"], c["site"]) if x]
    contact += [s.get("label") for s in c["social"] if s.get("label")]
    lines_section(t("Contact"), contact, fill="FFFFFF", color=INK, text_size=12.5)
    if r["summary"]:
        st = Stack(ctx, side, W)
        b = st.row()
        heads.append(pill(ctx, b, t("SUMMARY"), W, fill="FFFFFF", color=INK, size=16, before=20,
                          after=20, bg=NAVY).spacer)
        run(ctx, b.p(line=1.6, ind_start=6, ind_end=6), r["summary"], size=12, color="FFFFFF")
        st.done()
    if r["education"]:
        st = Stack(ctx, side, W)
        for i, ed in enumerate(r["education"]):
            b = st.row(pad_top=10 if i else 0)          # 10px between the cards
            if i == 0:
                heads.append(pill(ctx, b, t("EDUCATION"), W, fill=INK, color="FFFFFF", size=16,
                                  before=20, after=20, spacing=0.5, bg=NAVY).spacer)
            # a grey card: the stacked dates, then degree / school
            card = rounded_block(ctx, b, W, fill=CARD, bg=NAVY, radius=12,
                                 pad=(12, 14, 12, 14))
            dates = [x for x in (ed["start"], "—" if ed["start"] and ed["end"] else "",
                                 ed["end"]) if x]
            tbl = card.table([tw(44), tw(W - 28 - 44)])   # 32px of date + the 12px gap
            d, txt = tbl.rows[0].cells
            fmt_cell(ctx, d, pad=(0, 0, 0, 12))
            db = Box(ctx, d, 0)
            for x in dates:
                run(ctx, db.p(line=1.4), x, "bold", size=11.5, color=INK)
            db.finish()
            tb = Box(ctx, txt, 0)
            run(ctx, tb.p(line=1.4), ed["degree"], "bold", size=11.5, color=INK)
            for v in (ed["school"], f"{t('GPA')} {ed['gpa']}" if ed["gpa"] else ""):
                if v:
                    run(ctx, tb.p(line=1.4), v, size=11.5, color="666666")
            tb.finish()
            card.finish_round()
        st.done()
    lines_section(t("SOFTWARE"), r["tools"], fill=INK, color="FFFFFF", cv_section=False)
    lines_section(t("LANGUAGES"), [joined([lg["name"], lg["level"]], " — ")
                                    for lg in r["languages"]], fill=INK, color="FFFFFF")

    page.fit_side(heads)       # a full side column gives up gaps, never spills

    # ---- main column
    T = page.main_text_px
    top = page.main_row(pad_top=8, pad_x=(GAP + 6, PAD + 6))
    for part in r["name"].split(" ", 1):
        run(ctx, top.p(line=0.95), part, "name", size=52, color=INK, caps=True, spacing=-1)
    if r["title"]:
        run(ctx, top.p(before=8), r["title"], size=18, color="333333", spacing=1)
    if r["achievements"]:
        box = page.main_row(pad_top=16)
        n = len(r["achievements"])
        tbl = box.table([tw(T) // n] * n, borders={"top": (1, RULE), "bottom": (1, RULE)})
        for cell, a in zip(tbl.rows[0].cells, r["achievements"]):
            fmt_cell(ctx, cell, pad=(0, 6, 0, 6))
            cb = Box(ctx, cell, 0, pad_top=14, pad_bottom=14)
            run(ctx, cb.p(align="center", line=1.0), a["metric"], "metric", size=26,
                color=CORAL)
            run(ctx, cb.p(align="center", before=5), a["label"], "bold", size=9, color=MUTED,
                spacing=1, caps=True)
            cb.finish()

    H = "role" if ctx.rtl else "bold"
    node = [0]
    for j, job in enumerate(r["experience"]):
        box = page.main_row(pad_top=16 if j == 0 else 12)
        if j == 0:
            pill(ctx, box, t("PROFESSIONAL EXPERIENCE"), T, fill=INK, color="FFFFFF", size=15,
                 pad_y=12, radius=16, spacing=1, after=16, bg=FRAME)
        # date column (52px, end-aligned, stacked) | 16px | rail + content
        tbl = box.table([tw(6 + 52 + 16), tw(T - 74)])
        d, body = tbl.rows[0].cells
        fmt_cell(ctx, d, pad=(0, 6, 0, 16))
        fmt_cell(ctx, body, pad=(0, 16, 0, 6), borders={"start": (2, INK)})
        db = Box(ctx, d, 0)
        for x in [x for x in (job["start"], "—" if job["start"] and job["end"] else "",
                              job["end"]) if x]:
            run(ctx, db.p(align="end", line=1.4), x, size=11.5, color="333333")
        db.finish()
        bb = Box(ctx, body, 0)
        p = bb.p()
        node[0] += 1
        # the 10px black node centred on the 2px rail at the cell's start edge
        # (x from the content column's physical left; the rail is 16px before it)
        col_w = T - 74 - 22
        x = (col_w + 16 - 6) if ctx.rtl else (-16 - 6)
        vml_oval(ctx, p, x_pt=pt(x), y_pt=0, d_pt=pt(10), fill=INK, stroke=None,
                 n=100 + node[0])
        run(ctx, p, job["role"], H, size=14, color=INK)
        meta = joined([job["company"], job["location"]], " · ")
        if meta:
            run(ctx, bb.p(after=4), meta, size=12, color=MUTED)
        bullets(ctx, bb, job["bullets"], size=11.5, color="333333", line=1.5)
        bb.finish()

    if r["skills"]:
        skills = r["skills"]
        rows = (len(skills) + 2) // 3
        for rr in range(rows):
            box = page.main_row(pad_top=16 if rr == 0 else 0)
            if rr == 0:
                pill(ctx, box, t("CREATIVE & TECHNICAL SKILLS"), T, fill=INK, color="FFFFFF",
                     size=15, pad_y=12, radius=16, spacing=1, after=6, bg=FRAME)
            col = T / 3
            tbl = box.table([tw(col)] * 3)
            for k in range(3):
                cell = tbl.rows[0].cells[k]
                borders = {"end": (1, RULE), "bottom": (1, RULE)}
                if k == 0:
                    borders["start"] = (1, RULE)
                if rr == 0:
                    borders["top"] = (1, RULE)
                fmt_cell(ctx, cell, pad=(0, 18, 0, 18), borders=borders)
                i = rr * 3 + k
                cb = Box(ctx, cell, 0, pad_top=14, pad_bottom=14)
                if i < len(skills):
                    sk = skills[i]
                    # the name block is at least 34px tall (CSS min-height): a row
                    # with an AT LEAST height, Word's own equivalent
                    nt = cb.table([tw(col - 36)])
                    row_height(nt.rows[0], 34)
                    run(ctx, fmt_p(ctx, nt.rows[0].cells[0].paragraphs[0], line=1.3), sk["name"],
                        size=13, color=INK)
                    cb.pad_top = 16
                    if rated(sk):
                        pct, label = skill_pct(sk)
                        # min-width:52px in the PDF: a level word widens it
                        lw = max(52, text_px(ctx, "bold", label, 12) + 2)
                        sw = col - 36 - 10 - lw
                        inner = cb.table([tw(sw + 10), tw(lw)])
                        a, lab = inner.rows[0].cells
                        fmt_cell(ctx, a, valign="center")
                        fmt_cell(ctx, lab, valign="center")
                        slider_shape(ctx, a.paragraphs[0], pct, sw, on=CORAL, off=TRACK)
                        fmt_p(ctx, a.paragraphs[0], line=1.0)
                        run(ctx, fmt_p(ctx, lab.paragraphs[0], align="end"), label, "bold",
                            size=12, color=INK)
                cb.finish()

    if r["recognition"]:
        box = page.main_row(pad_top=16)
        pill(ctx, box, t("CERTIFICATIONS"), T, fill=INK, color="FFFFFF", size=15, pad_y=12,
             radius=16, spacing=1, after=16, bg=FRAME)
        for i, rec in enumerate(r["recognition"]):
            p = box.p(before=6 if i else 0, ind_start=6)
            run(ctx, p, rec["title"], "bold", size=12, color="333333")
            if rec["detail"]:
                run(ctx, p, " — " + rec["detail"], size=12, color="333333")
    page.finish()
