"""modern-t14 — Cream & Navy, orange flags: a cream column (0-366px) and a navy
column (366px-) down every page; across the top an orange panel (0-483px,
70px down) with a tracked Anton caps name, the title and a justified summary,
overlapping a navy-framed photo block; orange flags (overhanging the cream
column by 17px) over labelled contact, skills with dot rows, languages; navy
column: Anton heads over a grey hairline, education and experience on a thin
timeline with white nodes (title, place + dates at the end, a justified
paragraph), Recognition with a run-in Tools line. From modern/t14.j2.

Simplified in Word (overlapping positioned boxes): the header is its own
table - the orange panel a real cell, the photo beside it showing the part the
panel does not cover (483-761px)."""
from __future__ import annotations

from docx.shared import Pt

from ..docx_design import (
    Box, Ctx, design, dots_shape, fmt_cell, fmt_p, fmt_table, page_rects, photo_run, pt,
    row_height, run, tiny, tw, vmerge, vml_anchored, vml_oval,
)
from docx.oxml.ns import qn

from .common import SidebarPage, Stack, block, date_range, joined, single_px

CREAM, NAVY, ORANGE, PHOTO = "F7F4EC", "0E2B53", "B93E12", "2A4570"
INK, LVL, DOT_OFF, RULE = "2E2C28", "6F6A5E", "BEB9A6", "8A8A93"
SIDE_PX, PANEL_W, HEAD_PX = 366, 483, 346


@design("modern-t14")
def build(ctx: Ctx) -> None:
    r, t = ctx.r, ctx.t
    sec = ctx.doc.sections[0]
    sec.top_margin = Pt(0.1)
    sec.bottom_margin = Pt(pt(24))
    page_rects(ctx, [(0, SIDE_PX + 1, 0, 1100, CREAM), (SIDE_PX, 850 - SIDE_PX, 0, 1100, NAVY)])
    H = "role" if ctx.rtl else "body"

    # ---- the header table: [cream 70 / orange panel / cream] | [photo, merged]
    head = ctx.doc.add_table(rows=3, cols=2)
    fmt_table(head, [tw(PANEL_W), tw(850 - PANEL_W)])
    row_height(head.rows[0], 70)
    for i in range(3):
        vmerge(head.rows[i].cells[1], "restart" if i == 0 else "continue")
    pc = head.rows[0].cells[1]
    # the photo starts 30px down. Word drops a paragraph's space-before at
    # the top of the page (it sat at 0px), and a cell top margin would push
    # the whole row (Word gives every cell the row's largest), so the gap is
    # an empty line exactly as tall: 20pt mark, AUTO multiple of its line
    gap = pc.paragraphs[0]
    pp = pc.add_paragraph()
    fmt_p(ctx, pp, line=1.0)
    photo_run(ctx, pp, size_px=761 - PANEL_W, height_px=300, shape="rect",
              placeholder=PHOTO)
    panel = head.rows[1].cells[0]
    fmt_cell(ctx, panel, fill=ORANGE, pad=(0, 50, 0, 34))
    nb = Box(ctx, panel, 0, pad_top=22, pad_bottom=24)
    run(ctx, nb.p(line=1.1), r["name"], "name", size=38, color="FFFFFF", caps=True, spacing=5)
    if r["title"]:
        run(ctx, nb.p(before=2), r["title"], size=18, color="F6DDD3")
    if r["summary"]:
        run(ctx, nb.p(before=14, line=1.55, align="both"), r["summary"], size=12.5,
            color="FBE9E2")
    nb.finish()
    for cell in (head.rows[0].cells[0], head.rows[2].cells[0]):
        tiny(cell.paragraphs[0])
    for i in (1, 2):
        tiny(head.rows[i].cells[1].paragraphs[0])
    # after that loop: python-docx hands back the merge's TOP cell for those
    # continuation cells, so it re-tinied this very paragraph (photo at 0px)
    tiny(gap, 40)
    gap._p.find(qn("w:pPr") + "/" + qn("w:spacing")).set(
        qn("w:line"), str(round(240 * 30 / single_px(ctx, "body", 20 / 0.72))))
    # the gap down to the columns (346px main, 360px side): an at-least row
    # measured from the panel's foot is unknown, so the columns start right
    # after the photo (330px) + 16px
    # the columns start at 346px: the last header row makes up what the
    # 70px top row and the (measured) orange panel leave
    def size_head():
        from ..docx_measure import Measure
        from docx.oxml.ns import qn
        ph = Measure(ctx.resolved).block(panel._tc, pt(PANEL_W - 84))
        tr = head.rows[2]._tr
        for old in tr.findall(".//" + qn("w:trHeight")):
            old.getparent().remove(old)
        row_height(head.rows[2], max(16, HEAD_PX - 70 - ph / 0.72))
    ctx.after_lines.append(size_head)
    spacer = ctx.doc.add_paragraph()
    tiny(spacer)

    page = SidebarPage(ctx, SIDE_PX, side_fill=CREAM, side_pad=(14, 0, 24, 0),
                       main_pad_x=(44, 40), full_height_fill=False)
    side = page.side
    W = SIDE_PX - 78 - 24

    def flag(b, label):
        """383px orange flag from the page edge: 366px block + the 17px that
        overhang the navy column (an anchored shape, as t4's ribbon)."""
        # the PDF centres the label across the whole 383px flag (overhang
        # included): 17px more start padding centres it there (it sat 8.7px
        # short - measured, run 6)
        fb = block(ctx, b, SIDE_PX, fill=ORANGE, pad=(8, 39, 8, 0))
        p = fb.p(align="center")
        run(ctx, p, label, "heading", size=22, color="FFFFFF", caps=True, spacing=6)
        line_px = 22 * 1.5
        x = (SIDE_PX - 22) if not ctx.rtl else -17
        vml_anchored(p, x_pt=pt(x), y_pt=0, w_pt=pt(17), h_pt=pt(line_px + 16), fill=ORANGE)
        fb.finish()

    done = [0]

    def section(label, items, write):
        if not items:
            return
        st = Stack(ctx, side, SIDE_PX)
        for i, it in enumerate(items):
            b = st.row()
            if i == 0:
                # the column's 13px flex gap above every flag but the first
                # (measured: the 2nd and 3rd flags sat on the text above)
                if side.pending is None or done[0]:
                    b.pad_top = 13
                done[0] += 1
                flag(b, label)
                b.pad_top = 13
            write(b, i, it)
        st.done()

    def contact(b, i, item):
        label, value = item
        run(ctx, b.p(before=10 if i else 0, ind_start=78, ind_end=24), label, "bold", size=9.5,
            color=ORANGE, spacing=1.6, caps=True)
        run(ctx, b.p(ind_start=78, ind_end=24), value, size=12.5, color=INK)

    def skill(b, i, sk):
        # the level word ends 7px in from the column edge in the PDF (it sat
        # 17px short at SIDE_PX - 24 - measured, run 6)
        p = b.p(before=9 if i else 0, ind_start=78, tabs=[(SIDE_PX - 7, "end", None)])
        run(ctx, p, sk["name"], "bold", size=12.5, color=INK)
        if sk["level"]:
            run(ctx, p, "\t", size=11)
            run(ctx, p, sk["level"], size=11, color=LVL)
            dots_shape(ctx, b.p(before=5, ind_start=78), sk["dots"], sk["dot_total"], on=ORANGE,
                       off=DOT_OFF, d=9, gap=5)

    c = r["contact"]
    items = [(t("Phone"), c["phone"]), (t("Email"), c["email"]), (t("Studio"), c["address"]),
             (t("Portfolio"), c["site"])]
    section(t("Contact"), [x for x in items if x[1]], contact)
    section(t("Skills"), r["skills"], skill)
    section(t("Languages"), [joined([x["name"], x["level"]], " — ") for x in r["languages"]],
            lambda b, i, s: run(ctx, b.p(line=1.7, ind_start=78, ind_end=24), s, size=12.5,
                                color=INK))

    # ---- navy column
    T = page.main_text_px

    def nhead(box, label, role="heading"):
        run(ctx, box.p(), label, role, size=25, color="FFFFFF", spacing=1)
        rule = box.p()
        tiny(rule)
        # the first entry sits ~3px closer to the rule in the PDF (measured in
        # Word, run 6)
        fmt_p(ctx, rule, before=9, after=11, border={"bottom": (1, RULE, 0)})

    node = [0]

    def entry(box, title, place, dates, text):
        tbl = box.table([tw(T - 5)], ind=tw(5))
        cell = tbl.rows[0].cells[0]
        fmt_cell(ctx, cell, fill=NAVY, pad=(0, 26, 0, 0), borders={"start": (1, RULE)})
        eb = Box(ctx, cell, 0)
        p = eb.p(line=1.25)
        node[0] += 1
        cw = T - 5 - 26
        x = (cw + 26 - 5) if ctx.rtl else (-26 - 5)
        vml_oval(ctx, p, x_pt=pt(x), y_pt=pt(6), d_pt=pt(10), fill="FFFFFF", stroke=None,
                 n=500 + node[0])
        run(ctx, p, title, H, size=15, color="FFFFFF")
        if place or dates:
            q = eb.p(before=2, tabs=[(cw, "end", None)])
            run(ctx, q, place, size=12.5, color="C7C7CE")
            if dates:
                run(ctx, q, "\t", size=11.5)
                run(ctx, q, dates, size=11.5, color="A9A9B2")
        if text:
            run(ctx, eb.p(before=6, line=1.5, align="both"), text, size=12, color="CFCFD5")
        eb.finish()

    for j, ed in enumerate(r["education"]):
        box = page.main_row(fill=NAVY, pad_top=0 if j == 0 else 13)
        if j == 0:
            nhead(box, t("Education"))
        entry(box, ed["degree"], ed["school"], date_range(ed["start"], ed["end"]),
              " ".join(ed["bullets"]))
    for j, job in enumerate(r["experience"]):
        box = page.main_row(fill=NAVY, pad_top=16 if j == 0 else 13)
        if j == 0:
            nhead(box, t("Experience"))
        entry(box, job["role"], joined([job["company"], job["location"]], " · "),
              date_range(job["start"], job["end"]), " ".join(job["bullets"]))
    if r["recognition"] or r["tools"]:
        box = page.main_row(fill=NAVY, pad_top=11)    # 16 sat 5px low (run 6)
        # the label is not a .cv-section in the PDF (Arabic: the body face)
        nhead(box, t("Recognition") if r["recognition"] else t("Tools"),
              role="heading" if not ctx.rtl else "role")
        for i, rec in enumerate(r["recognition"]):
            p = box.p(line=1.55)
            run(ctx, p, rec["title"], "bold", size=12, color="FFFFFF")
            if rec["detail"]:
                run(ctx, p, " — " + rec["detail"], size=12, color="CFCFD5")
        if r["tools"]:
            p = box.p(before=6 if r["recognition"] else 0, line=1.55)
            if r["recognition"]:
                run(ctx, p, t("Tools"), "heading", size=12, color="FFFFFF")
                run(ctx, p, " — ", size=12, color="CFCFD5")
            run(ctx, p, " · ".join(r["tools"]), size=12, color="CFCFD5")
    tail = page.finish()
    fmt_cell(ctx, tail.c, fill=NAVY)
