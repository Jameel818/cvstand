"""modern-t1 — Open, editorial: a white page (64/60/56px margins), a two-line
caps name, the summary with one phrase on a lime marker, a stat band between
a black and a grey rule; then a main column (Experience, Recognition,
Education) beside a 236px side column ruled at its start edge (Contact,
Capabilities as name + level + dot row, Tools in two columns, Languages).
From modern/t1.j2."""
from __future__ import annotations

from docx.shared import Pt

from ..docx_design import (
    Box, Ctx, _rpr_put, design, dots_shape, fmt_cell, fmt_p, fmt_table, pt, run, tw,
)
from .common import SidebarPage, Stack, bullets, date_range, joined

INK, BODY, MUTED, RULE, DOT_OFF, MARK = "16161A", "3A3A42", "5C5C66", "D6D6DA", "B8B8BE", "D8F035"
PADX, SIDE_TXT = 60, 236


@design("modern-t1")
def build(ctx: Ctx) -> None:
    r, t = ctx.r, ctx.t
    sec = ctx.doc.sections[0]
    sec.top_margin = Pt(pt(64))
    sec.bottom_margin = Pt(pt(56))
    doc = ctx.doc
    H = "role" if ctx.rtl else "heading"

    def body_p(**kw):
        kw.setdefault("ind_start", PADX)
        kw.setdefault("ind_end", PADX)
        return fmt_p(ctx, doc.add_paragraph(), **kw)

    for part in r["name"].split(" ", 1):
        run(ctx, body_p(line=0.94), part, "name", size=64, color=INK, caps=True, spacing=-2.2)
    if r["summary"]:
        p = body_p(before=14, line=1.6, ind_end=850 - PADX - 560)
        hl = r.get("summary_highlight") or ""
        s = r["summary"]
        if hl and hl in s:
            a, b = s.split(hl, 1)
            if a:
                run(ctx, p, a, size=16, color=BODY)
            m = run(ctx, p, hl, size=16, color=BODY)
            _rpr_put(m._r.get_or_add_rPr(), "shd", val="clear", color="auto", fill=MARK)
            if b:
                run(ctx, p, b, size=16, color=BODY)
        else:
            run(ctx, p, s, size=16, color=BODY)
    if r["achievements"]:
        from ..docx_design import tiny
        gap = doc.add_paragraph()
        tiny(gap, before_px=21)
        n = len(r["achievements"])
        W = 850 - 2 * PADX
        tbl = doc.add_table(rows=1, cols=n)
        fmt_table(tbl, [tw(W) // n] * n, ind=tw(PADX),
                  borders={"top": (1, INK), "bottom": (1, RULE)})
        for k, (cell, a) in enumerate(zip(tbl.rows[0].cells, r["achievements"])):
            fmt_cell(ctx, cell, pad=(0, 0, 0, 18 if k < n - 1 else 0))
            cb = Box(ctx, cell, 0, pad_top=16, pad_bottom=16)
            run(ctx, cb.p(line=1.0), a["metric"], "metric", size=30, color=INK)
            run(ctx, cb.p(before=6), a["label"], "bold", size=9, color=MUTED, spacing=1.4,
                caps=True)
            cb.finish()
        gap2 = doc.add_paragraph()
        tiny(gap2)

    # ---- the two columns: main | 236px side ruled at its start
    page = SidebarPage(ctx, SIDE_TXT + 30 + PADX, side_fill=None, side_pad=(0, 30, 0, PADX),
                       main_pad_x=(PADX, 34), full_height_fill=False, side_end=True)
    side_cell = page.tbl.rows[0].cells[1]
    fmt_cell(ctx, side_cell, borders={"start": (1, RULE)})
    side = page.side
    W = SIDE_TXT

    heads = []

    def shead(b, label, first):
        hp = b.p(before=21 if first else 51, after=12)
        if not first:
            heads.append(hp)
        run(ctx, hp, label, "heading", size=15, color=INK, caps=True, spacing=2.2)

    first = [True]

    def section(label, items, write):
        if not items:
            return
        st = Stack(ctx, side, W)
        for i, it in enumerate(items):
            b = st.row()
            if i == 0:
                shead(b, label, first[0])
                first[0] = False
            write(b, i, it)
        st.done()

    c = r["contact"]
    contact = [x for x in (c["email"], c["phone"], c["address"], c["site"]) if x]
    contact += [s.get("label") for s in c["social"] if s.get("label")]
    section(t("Contact"), contact, lambda b, i, s: run(ctx, b.p(line=1.85), s, size=12.5,
                                                         color=BODY))

    def skill(b, i, sk):
        p = b.p(before=11 if i else 0, tabs=[(W, "end", None)])
        run(ctx, p, sk["name"], "bold", size=12.5, color=INK)
        if sk["level"]:
            run(ctx, p, "\t", size=10.5)
            run(ctx, p, sk["level"], size=10.5, color=MUTED)
            dots_shape(ctx, b.p(before=5), sk["dots"], sk["dot_total"], on=INK, off=DOT_OFF,
                       d=9, gap=5)
    section(t("Capabilities"), r["skills"], skill)

    if r["tools"]:
        tools = r["tools"]
        rows = [tools[i:i + 2] for i in range(0, len(tools), 2)]
        col = (W - 14) / 2

        def tool_row(b, i, pair):
            tbl = b.table([tw(col + 14), tw(col)])
            for k, name in enumerate(pair):
                cell = tbl.rows[0].cells[k]
                fmt_cell(ctx, cell, pad=(7 if i else 0, 0, 0, 14 if k == 0 else 0))
                cb = Box(ctx, cell, 0)
                run(ctx, cb.p(), name, size=12, color=BODY)
                cb.finish()
            if len(pair) == 1:
                Box(ctx, tbl.rows[0].cells[1], 0).finish()
        section(t("Tools"), rows, tool_row)
    section(t("Languages"), [joined([x["name"], x["level"]], " — ") for x in r["languages"]],
            lambda b, i, s: run(ctx, b.p(line=1.7), s, size=12.5, color=BODY))
    # a full side column gives up gaps, never spills. The Arabic estimate
    # reads this column ~20pt longer than Word draws it (measured on the user's 12pt
    # Markazi file, run 8: ink ends at 719pt, the estimate at 740pt): both
    # the gap fit and the type fit discount it, so the PDF's 51px section
    # gaps stay (they had shrunk to ~2/3 for nothing)
    # The squeeze also counts as overflow here, so the type fit (the PDF's
    # autofit) acts first and the gaps only give up what is left. Arabic
    # only: English's estimate is not long (its demo spilled with the bias)
    if ctx.rtl:
        ctx.side_bias_pt = 20.0
        ctx.squeeze_is_over = True
    page.fit_side(heads)

    # ---- main column
    T = page.main_text_px
    gap = [21]

    def head(box, label, after):
        run(ctx, box.p(after=after), label, "heading", size=15, color=INK, caps=True,
            spacing=2.2)

    for j, job in enumerate(r["experience"]):
        box = page.main_row(pad_top=gap[0] if j == 0 else 20)
        if j == 0:
            head(box, t("Experience"), 14)
        p = box.p(tabs=[(T, "end", None)])
        run(ctx, p, job["role"], H, size=16, color=INK)
        dates = date_range(job["start"], job["end"], " — ")
        if dates:
            run(ctx, p, "\t", size=11.5)
            run(ctx, p, dates, size=11.5, color="6A6A74")
        meta = joined([job["company"], job["location"]], " · ")
        if meta:
            run(ctx, box.p(before=3, after=8), meta, size=13, color=BODY)
        bullets(ctx, box, job["bullets"], size=12.5, color="2C2C34", line=1.6)
    gap[0] = 32 if r["experience"] else 21
    if r["recognition"]:
        box = page.main_row(pad_top=gap[0])
        head(box, t("Recognition"), 12)
        for i, rec in enumerate(r["recognition"]):
            p = box.p(before=8 if i else 0, line=1.5)
            run(ctx, p, rec["title"], "bold", size=12.5, color=INK)
            if rec["detail"]:
                run(ctx, p, " — " + rec["detail"], size=12.5, color=INK)
        gap[0] = 32
    for j, ed in enumerate(r["education"]):
        box = page.main_row(pad_top=gap[0] if j == 0 else 10)
        if j == 0:
            head(box, t("Education"), 12)
        run(ctx, box.p(line=1.5), ed["degree"], H, size=14, color=INK)
        sub = joined([ed["school"], ed["end"]], " · ")
        if sub:
            run(ctx, box.p(line=1.5), sub, size=12.5, color=MUTED)
        if ed["gpa"]:
            run(ctx, box.p(line=1.5), f"{t('GPA')} {ed['gpa']}", size=12.5, color=MUTED)
        bullets(ctx, box, ed["bullets"], size=12, color="2C2C34", line=1.55, before=4)
    page.finish()
