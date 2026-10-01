"""modern-t12 — Off-white editorial: a #F7F7F5 page; an Anton caps name with
the title end-aligned on its line; Anton heads with a 2px trailing rule; the
summary justified; stats centred between hairlines; then two pairs of
columns - Work Experience on a 2px rail with hollow nodes ("Role (dates)",
company, one paragraph) | Contact end-aligned with big Anton labels; and
Educational History on the same rail | Skills as a 2-up grid of RINGS (the
percent as text, Anton caps name, level). From modern/t12.j2."""
from __future__ import annotations

from docx.shared import Pt

from ..docx_design import (
    Box, Ctx, design, fmt_cell, page_rects, pt, rated, ring_anchored, rule_heading, run,
    skill_pct, tw, vml_oval,
)
from .common import ColumnPage, joined, single_px

PAPER, INK, BODY, MUTED, RULE, TRACK = "F7F7F5", "111111", "222222", "555555", "D5D5D2", "B8B8B8"
PAD, SIDE, GAP = 46, 210, 28


@design("modern-t12")
def build(ctx: Ctx) -> None:
    r, t = ctx.r, ctx.t
    sec = ctx.doc.sections[0]
    sec.top_margin = Pt(pt(PAD - 2))
    sec.bottom_margin = Pt(pt(PAD))
    page_rects(ctx, [(0, 850, 0, 1100, PAPER)])
    page = ColumnPage(ctx, pad_x=(PAD, PAD))
    T = page.main_text_px
    MAIN = T - SIDE - GAP
    H = "role" if ctx.rtl else "bold"

    # name | title on one line (baseline): a two-cell table
    top = page.main_row()
    tbl = top.table([tw(T - 260), tw(260)])
    a, b = tbl.rows[0].cells
    fmt_cell(ctx, a, valign="bottom")
    fmt_cell(ctx, b, valign="bottom")
    run(ctx, a.paragraphs[0], r["name"], "name", size=66, color=INK, caps=True, spacing=0.5)
    from ..docx_design import fmt_p
    fmt_p(ctx, a.paragraphs[0], line=0.9)
    if r["title"]:
        run(ctx, fmt_p(ctx, b.paragraphs[0], align="end", after=6), r["title"], "name",
            size=24, color=INK, caps=True, spacing=1)

    def ah(box, label, width, size=33, before=0, after=0):
        rule_heading(ctx, box, label, width_px=width, size=size, color=INK, rule=INK,
                     gap=16, rule_px=2, before=before, after=after)

    if r["summary"]:
        box = page.main_row(pad_top=20)
        ah(box, t("ABOUT ME"), T, size=32, after=20)
        run(ctx, box.p(line=1.75, align="both"), r["summary"], size=12.5, color=BODY)
    ach = [x for x in r["achievements"] if x["metric"]]
    if ach:
        box = page.main_row(pad_top=20)
        n = len(ach)
        st = box.table([tw(T) // n] * n, borders={"top": (1, RULE), "bottom": (1, RULE)})
        for cell, x in zip(st.rows[0].cells, ach):
            fmt_cell(ctx, cell, pad=(0, 6, 0, 6))
            cb = Box(ctx, cell, 0, pad_top=16, pad_bottom=16)
            run(ctx, cb.p(align="center", line=1.0), x["metric"], "metric", size=24, color=INK)
            run(ctx, cb.p(align="center", before=5), x["label"], "bold", size=9, color=MUTED,
                spacing=1, caps=True)
            cb.finish()

    node = [0]

    def rail(cell_box, entries, title_size, text_size):
        """A 2px rail at the start edge, content 22px in, hollow 15px nodes."""
        tb = cell_box.table([tw(MAIN)])
        cell = tb.rows[0].cells[0]
        fmt_cell(ctx, cell, pad=(0, 22, 0, 0), borders={"start": (2, INK)})
        eb = Box(ctx, cell, 0)
        for i, (title, sub, text) in enumerate(entries):
            p = eb.p(before=20 if i else 0)
            node[0] += 1
            cw = MAIN - 22
            x = (cw + 22 + 1 - 7.5) if ctx.rtl else (-22 - 1 - 7.5)
            vml_oval(ctx, p, x_pt=pt(x), y_pt=pt(-2), d_pt=pt(15), fill=PAPER, stroke=INK,
                     weight_pt=pt(2), n=600 + node[0])
            run(ctx, p, title, H, size=title_size, color=INK)
            if sub:
                run(ctx, eb.p(before=1), sub, size=11.5, color=MUTED)
            if text:
                run(ctx, eb.p(before=4, line=1.5), text, size=text_size, color="333333")
        eb.finish()

    def paren(x):
        d = joined([x["start"], x["end"]], "–")
        return f" ({d})" if d else ""

    # ---- work experience | contact
    # one block: a heading row and a content row (each head stays with its column)
    box = page.main_row(pad_top=20)
    pt_ = box.table([tw(MAIN + GAP), tw(SIDE)], rows=2)
    hl, hr = pt_.rows[0].cells
    fmt_cell(ctx, hl, pad=(0, 0, 0, GAP), valign="center")
    fmt_cell(ctx, hr, valign="center")
    hlb = Box(ctx, hl, 0)
    ah(hlb, t("WORK EXPERIENCE"), MAIN, after=12)
    hlb.finish()
    hrb = Box(ctx, hr, 0)
    run(ctx, hrb.p(align="end", line=1.0, after=12), t("CONTACT"), "heading", size=33,
        color=INK)
    hrb.finish()
    l, rr = pt_.rows[1].cells
    fmt_cell(ctx, l, pad=(0, 0, 0, GAP))
    lb = Box(ctx, l, 0)
    if r["experience"]:
        rail(lb, [(j["role"] + paren(j), joined([j["company"], j["location"]], " · "),
                   " ".join(j["bullets"])) for j in r["experience"]], 15, 11.5)
    lb.finish()
    rb = Box(ctx, rr, 0)
    c = r["contact"]
    k = 0
    if c["address"]:
        run(ctx, rb.p(align="end", line=1.5), c["address"], size=12, color=BODY)
        k += 1
    for label, value in ((t("PHONE"), c["phone"]), (t("EMAIL"), c["email"]), (t("WEB"), c["site"])):
        if value:
            run(ctx, rb.p(align="end", before=8 if k else 0, line=0.9), label, "heading", size=26,
                color=INK)
            run(ctx, rb.p(align="end", before=3), value, size=12, color=BODY)
            k += 1
    for s in c["social"]:
        if s.get("label"):
            run(ctx, rb.p(align="end", before=8), s["label"], size=12, color=BODY)
    rb.finish()

    # ---- educational history | skills (rings)
    box = page.main_row(pad_top=20)
    et = box.table([tw(MAIN + GAP), tw(SIDE)])
    l, rr = et.rows[0].cells
    fmt_cell(ctx, l, pad=(0, 0, 0, GAP))
    lb = Box(ctx, l, 0)
    # the PDF prints these two heads even over an empty section; Word does
    # not leave an empty heading (decision in the audit)
    if r["education"]:
        ah(lb, t("EDUCATIONAL HISTORY"), MAIN, after=12)
    if r["education"]:
        rail(lb, [(e["degree"] + paren(e), e["school"], " ".join(e["bullets"]))
                  for e in r["education"]], 14, 11)
    lb.finish()
    rb = Box(ctx, rr, 0)
    skills = r["skills"]
    if skills:
        run(ctx, rb.p(align="end", line=1.0, after=12), t("SKILLS"), "heading", size=33,
            color=INK)
    if skills:
        col = (SIDE - 6) / 2
        st = rb.table([tw(col + 6), tw(col)], rows=(len(skills) + 1) // 2)
        for i, sk in enumerate(skills):
            cell = st.rows[i // 2].cells[i % 2]
            fmt_cell(ctx, cell, pad=(0, 0, 0, 6 if i % 2 == 0 else 0))
            cb = Box(ctx, cell, 0, pad_top=20 if i >= 2 else 0)
            if rated(sk):
                pct, _ = skill_pct(sk)
                line = single_px(ctx, "bold", 13)
                pd = (70 - line) / 2
                p = cb.p(align="center", before=pd, after=pd + 7)
                run(ctx, p, f"{round(pct)}%", "bold", size=13, color=INK)
                ring_anchored(p, x_pt=pt((col - 70) / 2), y_pt=pt(20 if i >= 2 else 0), d_px=70,
                              width_px=10, pct=pct, on=INK, off=TRACK)
            run(ctx, cb.p(align="center", line=1.15), sk["name"], "heading", size=11, color=INK,
                caps=True, spacing=0.3)
            if sk["level"]:
                run(ctx, cb.p(align="center"), sk["level"], size=9, color="5F5F5F")
            cb.finish()
        if len(skills) % 2:
            Box(ctx, st.rows[-1].cells[1], 0).finish()
    rb.finish()
    page.finish()
