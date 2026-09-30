"""modern-t2 — Navy & Gold: navy sidebar with a gold-ringed round photo, gold
band with the name, labelled contact grid, stats between rules, an experience
timeline with hollow gold nodes, skill bars with the percent as text,
References. Measured from app/templates/resumes/modern/t2.j2."""
from __future__ import annotations

from ..docx_design import (
    Box, Ctx, design, fmt_cell, fmt_p, inline_bar_cells, photo_run, pt, rated, rule_heading, run,
    skill_pct, tw, vml_oval,
)
from .common import SidebarPage, Stack, date_range, joined

NAVY, GOLD, TRACK = "002366", "FFD700", "E0E0E0"
SIDE_TEXT, SIDE_MUTED = "DFE3EA", "C8CFDB"
GREY, INK, INK2, INK3 = "667085", "20242E", "4A5160", "333A47"
SIDE_PX = 298


def _side_head(ctx, box, label, before):
    rule_heading(ctx, box, label, width_px=SIDE_PX - 56, size=16.5, color="FFFFFF", rule=GOLD,
                 spacing=3, before=before, after=10)


def _main_head(ctx, box, label, width_px, before=0):
    rule_heading(ctx, box, label, width_px=width_px, size=19, color=NAVY, rule=GOLD,
                 spacing=3, gap=14, before=before, after=16)


@design("modern-t2")
def build(ctx: Ctx) -> None:
    r, t = ctx.r, ctx.t
    sec = ctx.doc.sections[0]
    from docx.shared import Pt
    sec.top_margin = Pt(pt(34))
    sec.bottom_margin = Pt(pt(30))
    page = SidebarPage(ctx, SIDE_PX, side_fill=NAVY, side_pad=(0, 28, 30, 28))
    side = page.side

    # ---- side column: each section a Stack (heading kept with its first entry)
    p = side.p(align="center", after=0)
    photo_run(ctx, p, size_px=150, ring_px=5, ring=GOLD, placeholder=TRACK)
    gap, W = 26, SIDE_PX - 56

    def section(label, items, write):
        if not items:
            return
        st = Stack(ctx, side, W)
        for i, it in enumerate(items):
            b = st.row()
            if i == 0:
                _side_head(ctx, b, label, gap)
            write(b, i, it)
        st.done()

    def summary(b, i, s):
        run(ctx, b.p(align="both", line=1.65), s, size=12.5, color=SIDE_TEXT)

    def education(b, i, ed):
        run(ctx, b.p(before=14 if i else 0, line=1.4), ed["degree"], "bold", size=13,
            color="FFFFFF")
        if ed["school"]:
            run(ctx, b.p(line=1.5), ed["school"], size=12, color=SIDE_MUTED)
        if ed["start"] or ed["end"]:
            run(ctx, b.p(line=1.5), date_range(ed["start"], ed["end"]), size=12, color=GOLD)

    def skill(b, i, sk):
        top = 11 if i else 0
        if not rated(sk):
            run(ctx, b.p(before=top), sk["name"], size=12, color="FFFFFF")
            return
        pct, label = skill_pct(sk)
        bar_w, lab_w, gapw = 74, 34, 8
        name_w = W - bar_w - lab_w - 2 * gapw
        fill = round(bar_w * pct / 100)
        widths = [tw(name_w + gapw), tw(fill), tw(bar_w - fill), tw(gapw + lab_w)]
        tbl = b.table([w if w > 0 else 1 for w in widths])
        c = tbl.rows[0].cells
        fmt_cell(ctx, c[0], valign="center", pad=(top, 0, 0, 0))
        fmt_p(ctx, c[0].paragraphs[0])
        run(ctx, c[0].paragraphs[0], sk["name"], size=12, color="FFFFFF")
        inline_bar_cells(ctx, c[1], c[2], on=GOLD, off=TRACK, height=7)
        for cell in (c[1], c[2]):
            fmt_cell(ctx, cell, pad=(top, 0, 0, 0))
        fmt_cell(ctx, c[3], valign="center", pad=(top, 0, 0, 0))
        lp = c[3].paragraphs[0]
        fmt_p(ctx, lp, align="end")
        run(ctx, lp, label, size=10, color=GOLD)

    def language(b, i, lg):
        p = b.p(before=6 if i else 0, tabs=[(W, "end", None)])
        run(ctx, p, lg["name"], size=12, color="FFFFFF")
        if lg["level"]:
            run(ctx, p, "	", size=12)
            run(ctx, p, lg["level"], size=12, color=GOLD)

    def cert(b, i, rec):
        run(ctx, b.p(before=10 if i else 0, line=1.4), rec["title"], "bold", size=12.5,
            color="FFFFFF")
        if rec["detail"]:
            run(ctx, b.p(line=1.5), rec["detail"], size=12, color=GOLD)

    section(t("ABOUT ME"), [r["summary"]] if r["summary"] else [], summary)
    section(t("EDUCATION"), r["education"], education)
    section(t("SKILLS"), r["skills"], skill)
    section(t("LANGUAGE"), r["languages"], language)
    section(t("CERTIFICATIONS"), r["recognition"], cert)


    # ---- main column
    text_px = page.main_text_px
    band = page.main_row(fill=GOLD, pad_top=26, pad_bottom=26)
    p = band.p(line=1.0)
    run(ctx, p, r["name"], "name", size=58, color=NAVY, caps=True, spacing=1)
    if r["title"]:
        run(ctx, band.p(before=8), r["title"], size=17, color=NAVY, spacing=4)

    c = r["contact"]
    cells = [(t("PHONE"), c["phone"]), (t("WEBSITE"), c["site"]), (t("EMAIL"), c["email"]),
             (t("LOCATION"), c["address"])]
    cells = [x for x in cells if x[1]]
    first_pad = 26
    if cells:
        box = page.main_row(pad_top=first_pad)
        first_pad = 22
        col = (text_px - 20) / 2
        tbl = box.table([tw(col + 20), tw(col)], rows=(len(cells) + 1) // 2)
        for i, (label, value) in enumerate(cells):
            cell = tbl.rows[i // 2].cells[i % 2]
            fmt_cell(ctx, cell, pad=(10 if i >= 2 else 0, 0, 0, 20 if i % 2 == 0 else 0))
            cb = Box(ctx, cell, 0)
            run(ctx, cb.p(), label, "bold", size=10, color=GREY, spacing=1.5)
            run(ctx, cb.p(), value, size=12.5, color=INK)
            cb.finish()
        if len(cells) % 2:
            Box(ctx, tbl.rows[-1].cells[1], 0).finish()

    if r["achievements"]:
        box = page.main_row(pad_top=first_pad)
        first_pad = 22
        n = len(r["achievements"])
        each = tw(text_px) // n
        tbl = box.table([each] * n, borders={"top": (1, TRACK), "bottom": (1, TRACK)})
        for cell, a in zip(tbl.rows[0].cells, r["achievements"]):
            fmt_cell(ctx, cell, pad=(16, 6, 16, 6))
            cb = Box(ctx, cell, 0)
            run(ctx, cb.p(align="center", line=1.0), a["metric"], "metric", size=26,
                color=NAVY)
            run(ctx, cb.p(align="center", before=5), a["label"], "bold", size=10, color=GREY,
                spacing=1)
            cb.finish()

    if r["experience"]:
        node_px, line_x = 32, 7         # the line 6-8px in, the content 32px in
        for j, job in enumerate(r["experience"]):
            box = page.main_row(pad_top=first_pad if j == 0 else 0)
            if j == 0:
                _main_head(ctx, box, t("EXPERIENCE"), text_px)
                first_pad = 22
            last = j == len(r["experience"]) - 1
            tbl = box.table([tw(line_x), tw(node_px - line_x), tw(text_px - node_px)])
            a, b, content = tbl.rows[0].cells
            fmt_cell(ctx, a, borders={"end": (2, GOLD)})
            cb = Box(ctx, content, 0)
            fmt_cell(ctx, content, pad=(0, 0, 0 if last else 18, 0))
            p = cb.p(tabs=[(text_px - node_px - 2, "end", None)])
            # the hollow node, centred on the line
            vml_oval(ctx, a.paragraphs[0], x_pt=pt(line_x) - pt(12) / 2 + 0.75
                     if not ctx.rtl else -pt(12) / 2 - 0.75,
                     y_pt=pt(2), d_pt=pt(12), fill="FFFFFF", stroke=GOLD, weight_pt=pt(2),
                     n=j + 1)
            run(ctx, p, job["role"], "role", size=15, color=NAVY)
            dates = date_range(job["start"], job["end"])
            if dates:
                run(ctx, p, "\t", size=12)
                run(ctx, p, dates, size=12, color=GREY, italic=True)
            meta = joined([job["company"], job["location"]], " · ")
            if meta:
                run(ctx, cb.p(before=3), meta, size=12, color=INK2)
            if job["bullets"]:
                run(ctx, cb.p(before=6, line=1.6), " ".join(job["bullets"]), size=12.5,
                    color=INK3)
            cb.finish()
            Box(ctx, a, 0).finish()
            Box(ctx, b, 0).finish()

    if r["references"]:
        box = page.main_row(pad_top=first_pad)
        _main_head(ctx, box, t("REFERENCES"), text_px)
        col = (text_px - 20) / 2
        refs = r["references"]
        tbl = box.table([tw(col + 20), tw(col)], rows=(len(refs) + 1) // 2)
        for i, ref in enumerate(refs):
            cell = tbl.rows[i // 2].cells[i % 2]
            fmt_cell(ctx, cell, pad=(20 if i >= 2 else 0, 0, 0, 20 if i % 2 == 0 else 0))
            cb = Box(ctx, cell, 0)
            run(ctx, cb.p(), ref["name"], "bold", size=14, color=NAVY)
            if ref["title"]:
                run(ctx, cb.p(before=2), ref["title"], size=12, color=INK2)
            for k, v in enumerate([x for x in (ref["phone"], ref["email"]) if x]):
                run(ctx, cb.p(before=4 if k == 0 else 0, line=1.5), v, size=12, color=INK3)
            cb.finish()
        if len(refs) % 2:
            Box(ctx, tbl.rows[-1].cells[1], 0).finish()
    page.finish()
