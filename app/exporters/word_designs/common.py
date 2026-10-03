"""Building blocks several designs share: the sidebar page (a full-height
side column + one unbreakable row per main-column block) and small entry
helpers. Everything in template px (docx_design.PX)."""
from __future__ import annotations

from ..docx_design import (
    Box, Ctx, PAGE_W_PX, cant_split, fmt_cell, fmt_table, page_rects, run, tw, vmerge,
)


class SidebarPage:
    """One table: [side | main]. The side cell is ONE cell merged down every
    row; each main block is its own row that cannot split (Word honours that,
    not keep-with-next, inside tables). A last, splittable row takes whatever
    the side column holds beyond the main column, so a long side column flows
    to the next page instead of being cut."""

    def __init__(self, ctx: Ctx, side_px: float, *, side_fill=None, side_pad=(34, 28, 30, 28),
                 main_pad_x=(40, 40), full_height_fill=True, side_end=False, top_px=0.0,
                 ind_px: float = 0.0):
        self.ctx = ctx
        # `ind_px`: the table starts that far in (an inset side column on a
        # framed page, modern-t5); the main column still ends at the page edge
        self.side_w = tw(side_px)
        self.main_w = tw(PAGE_W_PX - ind_px) - self.side_w
        self.side_px, self.main_px = side_px, PAGE_W_PX - ind_px - side_px
        self.side_end = side_end
        self.main_pad_x = main_pad_x
        self.side_fill = side_fill
        widths = [self.main_w, self.side_w] if side_end else [self.side_w, self.main_w]
        self.tbl = ctx.doc.add_table(rows=1, cols=2)
        fmt_table(self.tbl, widths, ind=tw(ind_px))
        self.rows = []
        side = self._side_cell(self.tbl.rows[0])
        vmerge(side, "restart")
        fmt_cell(ctx, side, fill=side_fill, pad=(0, side_pad[1], 0, side_pad[3]))
        # the column's bottom padding is the PAGE's bottom margin (every design
        # sets it equal): written again in the cell it counted twice once Word
        # stopped hiding the cell-end paragraph (Box.finish), and the side
        # column spilled an empty page 2 (t4 Arabic)
        self.side = Box(ctx, side, self.side_w - tw(side_pad[1] + side_pad[3]),
                        pad_top=side_pad[0], pad_bottom=0)
        self._first = True
        if full_height_fill and side_fill:
            x = PAGE_W_PX - side_px if side_end else 0
            page_rects(ctx, [(x, side_px + 1, 0, 1100, side_fill)])

    def _side_cell(self, row):
        return row.cells[1] if self.side_end else row.cells[0]

    def _main_cell(self, row):
        return row.cells[0] if self.side_end else row.cells[1]

    def main_row(self, *, pad_top=0, pad_bottom=0, fill=None, pad_x=None, split=False) -> Box:
        if self._first:
            row = self.tbl.rows[0]
            self._first = False
        else:
            row = self.tbl.add_row()
            # hold the cell first: once marked "continue", python-docx's
            # row.cells hands back the merge's TOP cell instead
            cont = self._side_cell(row)
            vmerge(cont, "continue")
            fmt_cell(self.ctx, cont, fill=self.side_fill)
            from ..docx_design import tiny
            tiny(cont.paragraphs[0])                   # structural, merged away
        if not split:
            cant_split(row)
        cell = self._main_cell(row)
        px = pad_x or self.main_pad_x
        fmt_cell(self.ctx, cell, fill=fill, pad=(0, px[0], 0, px[1]))
        box = Box(self.ctx, cell, self.main_w - tw(px[0] + px[1]), pad_top=pad_top,
                  pad_bottom=pad_bottom)
        self.rows.append(box)
        return box

    def shrink_side_photo(self, photo, *, min_px: float = 40, reserve_pt: float = 4.0) -> None:
        """The PDF's side column is a flex column whose photo block may SHRINK
        (flex-shrink) when the column is full: the photo gives up height,
        cropped (object-fit: cover), before anything else moves. Same here,
        measured once the line heights are final."""
        from ..docx_design import PAGE_H_PT, crop_picture_height, pt
        from ..docx_measure import Measure
        sec = self.ctx.doc.sections[0]
        avail = PAGE_H_PT - sec.top_margin.pt - sec.bottom_margin.pt - pt(self.side.pad_bottom) - 1
        tc, width = self.side.c._tc, self.side.w / 20

        def go():
            from docx.oxml.ns import qn
            m = Measure(self.ctx.resolved)
            over = m.block(tc, width) + reserve_pt - (avail - self._above(m))
            if over <= 0:
                return
            ext = photo._r.find(".//" + qn("wp:extent"))
            h = int(ext.get("cy")) / 12700
            crop_picture_height(photo, max(pt(min_px), h - over))
        self.ctx.after_lines.append(go)

    def center_side(self, *, reserve_pt: float | None = None) -> None:
        """The PDF's side column is `justify-content:center`: half the page's
        slack goes above its first paragraph (measured once the line heights
        are final); nothing moves when the column is full."""
        from docx.oxml.ns import qn
        from ..docx_design import PAGE_H_PT, pt, tiny
        from ..docx_measure import Measure, bump_before
        sec = self.ctx.doc.sections[0]
        if self.side.pending is not None:
            tiny(self.side.pending)
        avail = PAGE_H_PT - sec.top_margin.pt - sec.bottom_margin.pt
        tc, width = self.side.c._tc, self.side.w / 20
        if reserve_pt is None:
            reserve_pt = 30.0 if self.ctx.rtl else 6.0   # the Arabic estimate runs short

        def go():
            m = Measure(self.ctx.resolved)
            slack = avail - self._above(m) - reserve_pt - m.block(tc, width)
            first = next(tc.iter(qn("w:p")), None)
            if slack > 0 and first is not None:
                bump_before(first, slack / 2)
        self.ctx.after_lines.append(go)

    def fit_side(self, heads, *, reserve_pt: float = 12.0) -> None:
        """Only the safety half of spread_side: a side column too tall for
        the page gives up space before `heads` instead of spilling."""
        self.spread_side(heads, reserve_pt=reserve_pt, grow=False)

    def spread_side(self, heads, *, reserve_pt: float | None = None, grow: bool = True) -> None:
        """The side column as the PDF's `justify-content: space-between`:
        the page's slack shared out before each of `heads` (python-docx
        paragraphs). The page's bottom margin stands in for the column's
        bottom padding, so the side box's own is dropped. Call before
        finish()."""
        from ..docx_design import PAGE_H_PT, pt, tiny
        from ..docx_measure import Measure, spread
        if reserve_pt is None:
            reserve_pt = 24.0 if self.ctx.rtl else 4.0   # the Arabic estimate runs short
        sec = self.ctx.doc.sections[0]
        avail = PAGE_H_PT - sec.top_margin.pt - sec.bottom_margin.pt
        if grow:
            self.side.pad_bottom = 0
        else:
            # fit only: the column keeps its bottom padding, inside the page
            avail -= pt(self.side.pad_bottom) + 1
        if self.side.pending is not None:
            tiny(self.side.pending)      # what finish() will make of it
        heads = [h._p if hasattr(h, "_p") else h for h in heads]
        tc, width = self.side.c._tc, self.side.w / 20
        # measured once the line heights are final (docx_design.true_lines)
        def go():
            m = Measure(self.ctx.resolved)
            spread(m, tc, width, heads, avail - self._above(m), reserve_pt, grow=grow)
        self.ctx.after_lines.append(go)

    def pin_top(self, box: Box, y_px: float, *, min_gap_px: float = 0) -> None:
        """Start main-column row `box` at `y_px` from the top of the page (an
        absolutely placed flow in the PDF, e.g. below a band): the rows above
        are measured once the line heights are final and the gap made up as
        space before `box`'s first paragraph (at least `min_gap_px`)."""
        from docx.oxml.ns import qn
        from ..docx_design import pt
        from ..docx_measure import Measure, _set_before
        sec = self.ctx.doc.sections[0]

        def go():
            m = Measure(self.ctx.resolved)
            above = 0.0
            for b in self.rows:
                if b is box:
                    break
                above += self._row_h(m, b)
            first = next(box.c._tc.iter(qn("w:p")), None)
            if first is not None:
                _set_before(first, max(pt(min_gap_px), pt(y_px) - sec.top_margin.pt - above))
        self.ctx.after_lines.append(go)

    def _above(self, m) -> float:
        """Height of what the body holds BEFORE this page's table (a header
        card or band table, modern-t14/t21): the measured helpers' page starts
        below it."""
        from docx.oxml.ns import qn
        h = 0.0
        width = 612.0
        for el in self.ctx.doc.element.body:
            if el is self.tbl._tbl:
                break
            if el.tag == qn("w:p"):
                h += m.paragraph(el, width)
            elif el.tag == qn("w:tbl"):
                h += m.table(el)
        return h

    @staticmethod
    def _row_h(m, b) -> float:
        """A main row's height: its content, or its AT-LEAST height if taller
        (a band row: t7/t11/t18)."""
        from docx.oxml.ns import qn
        h = m.block(b.c._tc, b.w / 20)
        tr = b.c._tc.getparent()
        th = tr.find(qn("w:trPr") + "/" + qn("w:trHeight")) if tr is not None else None
        if th is not None:
            h = max(h, int(th.get(qn("w:val"))) / 20)
        return h

    def push_to_bottom(self, box: Box, *, bottom_px: float = 0, reserve_pt: float = 24.0) -> None:
        """The PDF's `margin-top:auto` on a main-column block: the page's slack
        (one page, measured once the line heights are final) goes above
        `box`, so it sits at the foot of the column, `bottom_px` above the
        bottom margin. Nothing moves when the column already fills the page."""
        from docx.oxml.ns import qn
        from ..docx_design import PAGE_H_PT, pt
        from ..docx_measure import Measure, bump_before
        sec = self.ctx.doc.sections[0]
        avail = PAGE_H_PT - sec.top_margin.pt - sec.bottom_margin.pt - pt(bottom_px)

        def go():
            m = Measure(self.ctx.resolved)
            used = sum(self._row_h(m, b) for b in self.rows) + self._above(m)
            slack = avail - reserve_pt - used
            first = next(box.c._tc.iter(qn("w:p")), None)
            if slack > 0 and first is not None:
                bump_before(first, slack)
        if not hasattr(self, "_pushes"):
            self._pushes = []
        self._pushes.append(go)

    @property
    def main_text_px(self) -> float:
        return self.main_px - sum(self.main_pad_x)

    def _fit_main(self, max_over_pt: float = 160.0, reserve_pt: float | None = None) -> None:
        """Shrink-only safety for a CV that ALMOST fits one page (the demo), as
        the app's own auto-fit compresses the PDF's vertical rhythm: when the
        main column's estimate exceeds the page by at most `max_over_pt`, the
        space before each row's first paragraph shrinks first; if that is not
        enough, ALL paragraph spacing in the main column scales down (never
        below 40 %, never the text). A long CV flows to page 2 untouched."""
        from docx.oxml.ns import qn
        from ..docx_design import PAGE_H_PT
        from ..docx_measure import Measure, _before, _set_before
        sec = self.ctx.doc.sections[0]
        if reserve_pt is None:
            # measured: the Arabic estimate runs 20-50pt short on a full page
            # (joined-form widths under-count Arabic line lengths a little)
            reserve_pt = 50.0 if self.ctx.rtl else 10.0
        avail = PAGE_H_PT - sec.top_margin.pt - sec.bottom_margin.pt - reserve_pt

        def go():
            m = Measure(self.ctx.resolved)
            used = sum(self._row_h(m, b) for b in self.rows) + self._above(m)
            over = used - avail
            if over <= 0 or over > max_over_pt:
                return
            firsts = [next(b.c._tc.iter(qn("w:p")), None) for b in self.rows]
            firsts = [f for f in firsts if f is not None]
            have = sum(_before(f) for f in firsts)
            if have >= over:
                k = 1 - over / have
                for f in firsts:
                    _set_before(f, _before(f) * k)
                return
            # every paragraph's spacing in the main column, proportionally
            sps = []
            for b in self.rows:
                for p_ in b.c._tc.iter(qn("w:p")):
                    sp = p_.find(qn("w:pPr") + "/" + qn("w:spacing"))
                    if sp is not None:
                        sps.append(sp)
            total = sum((int(sp.get(qn("w:before"), 0)) + int(sp.get(qn("w:after"), 0))) / 20
                        for sp in sps)
            if total <= 0:
                return
            k = max(0.4, 1 - over / total)
            for sp in sps:
                for a in ("before", "after"):
                    v = int(sp.get(qn(f"w:{a}"), 0))
                    if v:
                        sp.set(qn(f"w:{a}"), str(int(v * k)))
        self.ctx.after_lines.append(go)

    def finish(self):
        self._fit_main()
        # pushes run AFTER the fit: computed before it, the slack ignored what
        # the fit then took out, and the foot block stopped short of the foot
        # (modern-t21 Arabic, ~55px)
        self.ctx.after_lines.extend(getattr(self, "_pushes", []))
        tail = self.main_row(split=True)
        self.side.finish()
        for b in self.rows:
            b.finish()
        return tail


def joined(parts, sep):
    return sep.join(p for p in parts if p)


def date_range(a, b, dash=" – "):
    return f"{a}{dash if a and b else ''}{b}"


class Stack:
    """A column section as a one-column nested table, one row per entry, every
    row unbreakable, the heading in the FIRST entry's row: inside a cell that
    breaks across pages Word ignores keep-with-next, but it keeps a row that
    cannot split whole - so a heading never ends a page alone."""

    def __init__(self, ctx: Ctx, box: Box, width_px: float):
        self.ctx, self.box, self.w = ctx, box, tw(width_px)
        self.tbl = None
        self.boxes: list[Box] = []

    def row(self, pad_top: float = 0) -> Box:
        if self.tbl is None:
            self.tbl = self.box.table([self.w])
            row = self.tbl.rows[0]
        else:
            self.boxes[-1].finish()
            row = self.tbl.add_row()
        cant_split(row)
        cell = row.cells[0]
        fmt_cell(self.ctx, cell, pad=(pad_top, 0, 0, 0))
        b = Box(self.ctx, cell, self.w)
        self.boxes.append(b)
        return b

    def done(self):
        if self.boxes:
            self.boxes[-1].finish()


def block(ctx: Ctx, box: Box, width_px: float, *, fill=None, pad=(0, 0, 0, 0), border=None,
          valign=None) -> Box:
    """A filled / bordered block (ribbon, pill, card, bar): a one-cell table.
    `pad` (top, start, bottom, end) px - top/bottom as paragraph spacing."""
    tbl = box.table([tw(width_px)])
    cell = tbl.rows[0].cells[0]
    fmt_cell(ctx, cell, fill=fill, pad=(0, pad[1], 0, pad[3]), borders=border, valign=valign)
    return Box(ctx, cell, tw(width_px - pad[1] - pad[3]), pad_top=pad[0], pad_bottom=pad[2])


def bullets(ctx: Ctx, box: Box, items, *, size, color, line=1.5, indent=16, bullet="•",
            bullet_color=None, before=0, gap=0, ind_start=0):
    """Real text bullets (a bullet, a tab, the text, a hanging indent), so the
    bullet takes the template's colour and size."""
    for i, b in enumerate(items):
        p = box.p(before=before if i == 0 else gap, line=line, ind_start=ind_start + indent,
                  hanging=indent, tabs=[(ind_start + indent, "start", None)])
        run(ctx, p, bullet + "	", size=size, color=bullet_color or color)
        run(ctx, p, b, size=size, color=color)


def node_x(ctx: Ctx, cell_px: float, at: str, d_px: float) -> float:
    """x (pt, from the cell's physical left) that centres a `d_px` node on the
    cell's LOGICAL `at` edge ("start"/"end") - a VML shape is placed physically."""
    from ..docx_design import pt
    left = (at == "start") != ctx.rtl
    return (0.0 if left else pt(cell_px)) - pt(d_px) / 2


_ROLE_FACE = {"heading": "heading", "name": "name", "bold": "body_bold", "role": "role",
              "body": "body", "metric": "heading", "accent": "body"}


def single_px(ctx: Ctx, role: str, size_px: float) -> float:
    """Word's single line (px) for `role`'s face at `size_px` (measured rule,
    docx_measure)."""
    from ..docx_measure import _face_file, _font
    key = _ROLE_FACE[role]
    if key == "heading" and role == "metric" and ctx.rtl:
        key = "role"
    f = ctx.resolved["faces"][key]
    rel = _face_file(f["family"], f["weight"])
    return size_px * (_font(rel)["line"] if rel else 1.2)


_CORNER = {  # VML paths of a corner MASK (square minus a quarter circle), r x r
    "tl": "m0,0 l{r},0 qx0,{r} x e",
    "tr": "m0,0 qx{r},{r} l{r},0 x e",
    "bl": "m0,0 qy{r},{r} l0,{r} x e",
    "br": "m{r},0 qy0,{r} l{r},{r} x e",
}


def round_corners(ctx: Ctx, inner: Box, width_px: float, radius_px: float, bg: str,
                  pad_left_px: float) -> None:
    """Round a filled block's corners: four tiny in-front shapes in the colour
    BEHIND the block mask its square corners (Word shading has none, and a
    shape behind text vanishes under a cell's shading). All four hang on the
    cell's FIRST paragraph; the block's height is measured once the line
    heights are final, and the radius never exceeds half the height or width
    (a pill). `pad_left_px`: the cell's physical left margin (masks are
    placed from the text column)."""
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph
    from ..docx_design import pt, vml_anchored
    from ..docx_measure import Measure
    tc = inner.c._tc
    first_el = tc.find(qn("w:p"))
    if first_el is None:
        first_el = next(tc.iter(qn("w:p")), None)
    if first_el is None:
        return
    first = Paragraph(first_el, inner.c)

    def go():
        h_px = Measure(ctx.resolved).block(tc, pt(width_px)) / 0.72
        r = max(1.0, min(radius_px, h_px / 2, width_px / 2))
        k = 100
        x_l, x_r = pt(-pad_left_px), pt(width_px - pad_left_px - r)
        # The bottom masks: from the block's TOP they rest on the estimated
        # height, and a block Word sets a few px taller (modern-t5's education
        # cards: a square strip showed under the rounded corners, run 6) was
        # left with its corners rounded above its real bottom. When the block
        # ends in an empty spacer paragraph after a table (its bottom padding,
        # as space before), they hang on THAT paragraph instead: its own
        # spacing is known exactly, nothing above is estimated.
        last_el = tc.findall(qn("w:p"))[-1]
        prev = last_el.getprevious()
        bottom_on, y_b = first, pt(h_px - r)
        if (last_el is not first_el and prev is not None and prev.tag == qn("w:tbl")
                and not "".join(t.text or "" for t in last_el.iter(qn("w:t"))).strip()):
            sp = last_el.find(qn("w:pPr") + "/" + qn("w:spacing"))
            before = int(sp.get(qn("w:before"), 0)) / 20 if sp is not None else 0.0
            bottom_on, y_b = Paragraph(last_el, inner.c), max(0.0, before + 0.3 - pt(r))
        for key, on, x, y in (("tl", first, x_l, 0.0), ("tr", first, x_r, 0.0),
                              ("bl", bottom_on, x_l, y_b), ("br", bottom_on, x_r, y_b)):
            vml_anchored(on, x_pt=x, y_pt=y, w_pt=pt(r), h_pt=pt(r), fill=bg,
                         path=_CORNER[key].format(r=k), coords=f"{k},{k}", z=20)
    ctx.after_lines.append(go)


def rounded_block(ctx: Ctx, box: Box, width_px: float, *, fill: str, bg: str, radius: float,
                  pad=(0, 0, 0, 0)) -> Box:
    """A filled block (one-cell table, as `block`) whose corners are rounded
    by round_corners() when its Box is finished through `.finish_round()`."""
    inner = block(ctx, box, width_px, fill=fill, pad=pad)
    if not inner.pad_top:
        # the corner masks hang on the cell's first paragraph: make sure there
        # is one ABOVE any table (a cell-end paragraph after a table is not
        # laid out by Word)
        inner.pad_top = 0.01
    pad_left = pad[3] if ctx.rtl else pad[1]

    def finish_round():
        inner.finish()
        round_corners(ctx, inner, width_px, radius, bg, pad_left)
    inner.finish_round = finish_round
    return inner


def pill(ctx: Ctx, box: Box, label: str, width_px: float, *, fill: str, color: str, size,
         pad_y=10, pad_x=22, radius=14, spacing=0, before=0, role="heading", after=0,
         align="start", bg="FFFFFF"):
    """A rounded pill heading: a filled one-cell block (the label live text in
    its role style) with its corners rounded against `bg`; `before`/`after`
    px around it."""
    spacer = None
    if before:
        # its own spacer paragraph: a column too full can shrink it (fit_side)
        from ..docx_design import tiny
        spacer = box.p()
        tiny(spacer, before_px=before + box.pad_top)
        box.pad_top = 0
    inner = rounded_block(ctx, box, width_px, fill=fill, bg=bg, radius=radius,
                          pad=(pad_y, pad_x, pad_y, pad_x))
    p = inner.p(align=align)
    run(ctx, p, label, role, size=size, color=color, spacing=spacing)
    inner.finish_round()
    if after:
        box.pad_top += after
    p.spacer = spacer
    return p


class ColumnPage(SidebarPage):
    """One full-width column (a band template with no side column): one
    unbreakable row per block, the same measured helpers (pin_top,
    push_to_bottom, the near-one-page fit) as SidebarPage."""

    def __init__(self, ctx: Ctx, *, pad_x=(48, 48)):
        self.ctx = ctx
        self.side = None
        self.side_end = False
        self.main_pad_x = pad_x
        self.side_fill = None
        self.side_px, self.main_px = 0, PAGE_W_PX
        self.side_w, self.main_w = 0, tw(PAGE_W_PX)
        self.tbl = ctx.doc.add_table(rows=1, cols=1)
        fmt_table(self.tbl, [self.main_w])
        self.rows = []
        self._first = True

    def _main_cell(self, row):
        return row.cells[0]

    def main_row(self, *, pad_top=0, pad_bottom=0, fill=None, pad_x=None, split=False) -> Box:
        if self._first:
            row = self.tbl.rows[0]
            self._first = False
        else:
            row = self.tbl.add_row()
        if not split:
            cant_split(row)
        cell = row.cells[0]
        px = pad_x or self.main_pad_x
        fmt_cell(self.ctx, cell, fill=fill, pad=(0, px[0], 0, px[1]))
        box = Box(self.ctx, cell, self.main_w - tw(px[0] + px[1]), pad_top=pad_top,
                  pad_bottom=pad_bottom)
        self.rows.append(box)
        return box

    def finish(self):
        self._fit_main()
        self.ctx.after_lines.extend(getattr(self, "_pushes", []))   # after the fit
        tail = self.main_row(split=True)
        for b in self.rows:
            b.finish()
        return tail


def text_px(ctx: Ctx, role: str, text: str, size_px: float, spacing_px: float = 0.0) -> float:
    """Width (px) of `text` in `role`'s face (font advance widths, joined
    Arabic forms) - for shrink-wrapped boxes (an inline-block pill)."""
    from ..docx_measure import _face_file, _font, _w
    key = _ROLE_FACE[role]
    f = ctx.resolved["faces"][key]
    rel = _face_file(f["family"], f["weight"])
    if not rel:
        return len(text) * size_px * 0.55
    m = _font(rel)
    return sum(_w(m, ch) for ch in text) * size_px + spacing_px * len(text)
