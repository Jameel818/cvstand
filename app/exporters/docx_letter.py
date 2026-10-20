"""The cover letter in Word: one page, editable, in the template's identity.

The same identity as the PDF letter (app/letter.py: letter_themes.json), drawn
with Word's own editable objects - paragraphs, a shaded table cell for the
band, a header rectangle for the side rail - and the same role styles as the
CV (CVName, CVHeading, Normal), so the fonts follow the template and the font
ribbon exactly as the CV's do (docx_theme.resolve / apply), and are embedded
the same way (docx_font_embed).

Two entry points:
    render(data, key)          the letter alone, a new document
    append(doc, data, key)     the letter as its OWN SECTION after a finished
                               CV document ("CV + cover letter"): a new page
                               with its own (empty) header, so none of the
                               CV's page shapes - a sidebar fill, a band -
                               reach the letter's page.

Arabic: w:bidi on the letter's paragraphs and tables and the section, w:rtl
only on runs holding Arabic letters (docx_design.run), as the CV designs do.
"""
from __future__ import annotations

import io

from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Twips

from .. import letter as letter_mod
from ..schema import lang_of, normalize
from . import docx_design as D
from . import docx_font_embed, docx_theme
from .docx_layout import PPR_ORDER, TBLPR_ORDER, _header_shapes, _put, _w

PAGE_H_PX = 1100
MARGIN_PX = 68
RAIL_PX = 40


#: Word has no auto-fit, and its lines for the Latin body faces run taller
#: than the PDF's (each face's own single line - docx_design.true_lines), so
#: a long ENGLISH letter takes its body a step smaller in Word than in the
#: PDF. Measured in real Word (run 10): at the PDF's sizes the longest letter
#: the form allows spilled 38-151pt onto a second page in all 10 templates
#: checked, while Arabic, whose faces wrap fewer characters per line, ended
#: ~60pt clear. (characters of greeting + paragraphs + closing, body px)
WORD_BODY_SIZES_EN = ((1400, 13.0), (1800, 12.0), (2200, 11.0), (2600, 10.5), (10 ** 9, 10.0))


def word_body_px(letter: dict, lang: str) -> float:
    if lang == "ar":
        return letter_mod.body_px(letter)
    n = len(letter["greeting"]) + len(letter["closing"]) + sum(len(p) for p in letter["paragraphs"])
    return next(px for limit, px in WORD_BODY_SIZES_EN if n <= limit)


def _hex(c: str | None) -> str | None:
    return c.lstrip("#").upper() if c else None


def _section(doc, *, new: bool, rtl: bool, th: dict):
    """The letter's page: Letter size, its margins, an empty header of its
    own (flat, so it takes no room) and the side rail as a page shape."""
    if new:
        sec = doc.add_section(WD_SECTION.NEW_PAGE)
        # the paragraph that now ends the CV's section must add nothing. Not
        # D.tiny(): it APPENDS w:rPr, which here would land after w:sectPr -
        # schema-invalid, and Word refuses the whole file (docx gotcha 0).
        ppr = D._ppr(doc.element.body.findall(qn("w:p"))[-1])
        rpr = OxmlElement("w:rPr")
        rpr.append(_w("sz", val=2))
        rpr.append(_w("szCs", val=2))
        _put(ppr, rpr, PPR_ORDER)
        _put(ppr, _w("spacing", before=0, after=0, line=60, lineRule="auto"), PPR_ORDER)
    else:
        sec = doc.sections[0]
    sec.different_first_page_header_footer = False
    sec.page_width, sec.page_height = Twips(12240), Twips(15840)
    rail = th["rail"]
    start = MARGIN_PX + (RAIL_PX if rail and rail["at"] == "start" else 0)
    end = MARGIN_PX + (RAIL_PX if rail and rail["at"] == "end" else 0)
    left, right = (start, end) if not rtl else (end, start)
    sec.left_margin, sec.right_margin = Twips(D.tw(left)), Twips(D.tw(right))
    sec.top_margin = Twips(0 if th["head_bg"] else D.tw(56))
    sec.bottom_margin = Twips(D.tw(40))
    sec.header_distance = sec.footer_distance = 0
    for part in (sec.header, sec.footer):
        part.is_linked_to_previous = False
        for para in part.paragraphs:
            D.tiny(para)
    if rail:
        at_left = (rail["at"] == "start") != rtl
        x = 0 if at_left else D.PAGE_W_PX - RAIL_PX
        _header_shapes(sec, [(D.pt(x), D.pt(RAIL_PX), 0, D.pt(PAGE_H_PX), _hex(rail["bg"]))], [])
    return sec, start, end


def _build(doc, data: dict, key: str, resolved: dict, *, new_section: bool) -> None:
    from ..labels import reset_lang, set_lang, t

    lang = lang_of(data)
    rtl = lang == "ar"
    th = letter_mod.theme(key, lang)
    L = letter_mod.letter_of(data)
    r = normalize(data)
    body = doc.element.body
    existing = set(body)            # held, so lxml keeps the same proxies
    band = []
    _, start, end = _section(doc, new=new_section, rtl=rtl, th=th)
    token = set_lang(lang)
    try:
        ctx = D.Ctx(doc=doc, r=r, lang=lang, resolved=resolved, photo=None, t=t)
        page = D.Box(ctx, doc, D.tw(D.PAGE_W_PX - start - end))
        align = "center" if th["centred"] else "start"
        head_bg = _hex(th["head_bg"])

        def head(box: D.Box) -> None:
            p = box.p(align=align, line=1.1)
            D.run(ctx, p, r["name"], "name", size=36, color=_hex(th["name"]),
                  caps=th["name_caps"], spacing=th["name_tracking"] or None)
            if r["title"]:
                caps = th["heading_caps"]
                p = box.p(align=align, before=9, line=1.3)
                D.run(ctx, p, r["title"], "body", size=12.5 if caps else 15, bold=True,
                      color=_hex(th["title"]), caps=caps, spacing=1.6 if caps else None)
            items = letter_mod.contact_items(r)
            if items:
                p = box.p(align=align, before=14, line=1.45)
                D.run(ctx, p, D._SEP.join(items), "body", size=11.5, color=_hex(th["contact"]))

        if head_bg:
            band_w = D.PAGE_W_PX - (RAIL_PX if th["rail"] else 0)
            tbl = page.table([D.tw(band_w)], rows=2, ind=-D.tw(MARGIN_PX))
            band.append(tbl._tbl)
            cell = D.fmt_cell(ctx, tbl.rows[0].cells[0], fill=head_bg,
                              pad=(0, MARGIN_PX, 0, MARGIN_PX))
            box = D.Box(ctx, cell, D.tw(band_w - 2 * MARGIN_PX), pad_top=46, pad_bottom=34)
            head(box)
            box.finish()
            stripe = D.fmt_cell(ctx, tbl.rows[1].cells[0], fill=_hex(th["stripe"]))
            D.row_height(tbl.rows[1], 5, "exact")
            D.tiny(stripe.paragraphs[0])
        else:
            head(page)
            rule = page.table([D.tw(56)], rows=1, jc="center" if th["centred"] else None)
            cell = D.fmt_cell(ctx, rule.rows[0].cells[0], fill=_hex(th["accent"]))
            D.row_height(rule.rows[0], 3, "exact")
            D.tiny(cell.paragraphs[0])
            # the rule sits under the header with the PDF's 26px above it
            prev = rule._tbl.getprevious()
            if prev is not None and prev.tag == qn("w:p"):
                sp = prev.find(qn("w:pPr")).find(qn("w:spacing"))
                sp.set(qn("w:after"), str(int(round(D.pt(26) * 20))))

        size = word_body_px(L, lang)
        text = _hex(th["text"])
        gap = 34
        if L["date"]:
            D.run(ctx, page.p(align="end", before=gap, line=1.62), L["date"], "body",
                  size=size, color=text)
            gap = 18
        for i, line in enumerate(letter_mod.recipient_lines(L)):
            D.run(ctx, page.p(before=gap if i == 0 else 0, line=1.62), line, "body", size=size,
                  color=text, bold=True if i == 0 else None)
            gap = 26
        if L["subject"]:
            p = page.p(before=gap, line=1.35, keep=True)
            D.run(ctx, p, "■ ", "body", size=11, color=_hex(th["accent"]))
            D.run(ctx, p, L["subject"], "heading", size=14, color=_hex(th["heading"]),
                  caps=th["heading_caps"], spacing=1.2 if th["heading_caps"] else None)
            gap = 22
        if L["greeting"]:
            D.run(ctx, page.p(before=gap, line=1.62), L["greeting"], "body", size=size,
                  color=text)
            gap = 12
        for para in L["paragraphs"]:
            D.run(ctx, page.p(before=gap, line=1.62), para, "body", size=size, color=text)
            gap = 12
        if L["closing"]:
            D.run(ctx, page.p(before=24, line=1.62, keep=True), L["closing"], "body",
                  size=size, color=text)
        if L["signature"]:
            D.run(ctx, page.p(before=8 if L["closing"] else 24, line=1.2), L["signature"],
                  "name", size=21, color=_hex(th["heading"]))
        D.true_lines(ctx)
    finally:
        reset_lang(token)
    D.end_paragraph(doc)
    if rtl:
        # Only what the letter added - the CV before it is already finished.
        # Told apart by IDENTITY, not position: new content goes in BEFORE the
        # body's closing sectPr, so "everything after the old length" skipped
        # the letter's first element (its band, or its name line), which then
        # kept left-to-right properties (found in real Word, run 10).
        for el in [e for e in body if e not in existing]:
            for p in el.iter(qn("w:p")):
                _put(D._ppr(p), _w("bidi"), PPR_ORDER)
            for tbl in el.iter(qn("w:tbl")):
                if tbl in band:
                    # One cell, nothing to mirror - and Word applies the band's
                    # negative indent from the wrong side in a bidiVisual
                    # table: the band slid 49pt off the right edge, cutting
                    # the name (measured in Word, run 10).
                    continue
                _put(tbl.find(qn("w:tblPr")), _w("bidiVisual"), TBLPR_ORDER)
        _section_bidi(doc.sections[-1]._sectPr)


def _section_bidi(sect) -> None:
    """w:bidi on a section, at its schema position (as docx_design.finish_rtl)."""
    if sect.find(qn("w:bidi")) is not None:
        return
    el = OxmlElement("w:bidi")
    later = ("rtlGutter", "docGrid", "printerSettings", "sectPrChange")
    for i, ch in enumerate(sect):
        if ch.tag.rsplit("}", 1)[-1] in later:
            sect.insert(i, el)
            return
    sect.append(el)


# ---- a Modern template's letter: its own Word DESIGN, the CV content out ------------

#: what the letter takes out of the CV (the letterhead - name, headline,
#: contact, photo - stays, drawn by the template's own design)
CV_CONTENT = ("summary", "summary_highlight", "achievements", "experience", "recognition",
              "education", "skills", "tools", "languages", "references")


def _emptied(data: dict) -> dict:
    import copy
    d = copy.deepcopy(data)
    for k in CV_CONTENT:
        d[k] = "" if k.startswith("summary") else []
    return d


def _run_colour(doc, text: str) -> str | None:
    """The colour of the run that draws `text` (the CV's first bullet)."""
    if not text:
        return None
    probe = text[:30]
    for r in doc.element.body.iter(qn("w:r")):
        if probe in "".join(t.text or "" for t in r.iter(qn("w:t"))):
            c = r.find(qn("w:rPr") + "/" + qn("w:color"))
            return c.get(qn("w:val")) if c is not None else None
    return None


def _grid_cells(tr):
    """[(first grid column, tc)] of a row (gridSpan counted)."""
    out, col = [], 0
    for tc in tr.findall(qn("w:tc")):
        out.append((col, tc))
        pr = tc.find(qn("w:tcPr"))
        span = pr.find(qn("w:gridSpan")) if pr is not None else None
        col += int(span.get(qn("w:val"))) if span is not None else 1
    return out


def _job_column(full, text: str):
    """(table index, grid column) of the cell drawing `text` (the CV's first
    company) in the full CV design, or None."""
    if not text:
        return None
    tbls = [el for el in full.element.body if el.tag == qn("w:tbl")]
    for ti, tbl in enumerate(tbls):
        for tr in tbl.findall(qn("w:tr")):
            for col, tc in _grid_cells(tr):
                if text in "".join(t.text or "" for t in tc.iter(qn("w:t"))):
                    return ti, col
    return None


def _main_cell(doc, where=None):
    """(python-docx cell or None, text width in twips): the design's main
    column. `where` = (table, grid column) of the CV's jobs in the full
    design: the same column here - its merged first cell when the column runs
    down the page (modern-t14's main column IS the merged one), else its last
    row's cell. Without it: the widest cell of the layout table's last row
    that is not the merged side column. None for a full-width page."""
    from docx.table import Table, _Cell
    tbls = [el for el in doc.element.body if el.tag == qn("w:tbl")]
    if not tbls:
        return None, None
    if where is not None and where[0] < len(tbls):
        tbl = tbls[where[0]]
        rows = tbl.findall(qn("w:tr"))
        for tr in [rows[0]] + rows[::-1]:
            for col, tc in _grid_cells(tr):
                pr = tc.find(qn("w:tcPr"))
                vm = pr.find(qn("w:vMerge")) if pr is not None else None
                if col != where[1]:
                    continue
                if vm is not None and vm.get(qn("w:val")) != "restart":
                    continue                          # a continued merge: its first cell holds it
                if tr is rows[0] and vm is None and len(rows) > 1:
                    continue                          # not merged: take the last row's
                return _Cell(tc, Table(tbl, doc._body)), int(D._cell_w(pr) * 20)
    # the page's layout table: the one with the merged side column, else the
    # last (modern-t21 draws its header card as a table of its own first)
    sided = [t for t in tbls if t.find(".//" + qn("w:vMerge")) is not None]
    tbl = sided[0] if sided else tbls[-1]
    best = None
    for tc in tbl.findall(qn("w:tr"))[-1].findall(qn("w:tc")):
        pr = tc.find(qn("w:tcPr"))
        if pr is not None and pr.find(qn("w:vMerge")) is not None:
            continue
        w = D._cell_w(pr) if pr is not None else 0
        if best is None or w > best[1]:
            best = (tc, w)
    if best is None:
        return None, None
    return _Cell(best[0], Table(tbl, doc._body)), int(best[1] * 20)


def _drop_orphan_labels(doc) -> None:
    """A design that prints a section heading even over nothing (modern-t12's
    "WORK EXPERIENCE"): the letter has no sections, so no section names."""
    from .. import letter as letter_mod
    for p in list(doc.element.body.iter(qn("w:p"))):
        text = "".join(t.text or "" for t in p.iter(qn("w:t"))).strip()
        if text and letter_mod._is_label(text) and not letter_mod._is_label_contact(text):
            # the whole line: its words and the rule drawn after them (a tab
            # leader, modern-t12); the empty paragraph keeps the layout
            for r in p.findall(qn("w:r")):
                p.remove(r)


def _write_body(ctx, box, L: dict, lang: str, colour: str | None) -> None:
    """The plain letter body: date, recipient, subject, greeting,
    paragraphs, closing, signature - in the body face and colour."""
    size = word_body_px(L, lang)
    gap = 30
    if L["date"]:
        D.run(ctx, box.p(align="end", before=gap, line=1.62), L["date"], "body", size=size,
              color=colour)
        gap = 18
    for i, line in enumerate(letter_mod.recipient_lines(L)):
        D.run(ctx, box.p(before=gap if i == 0 else 0, line=1.62), line, "body", size=size,
              color=colour, bold=True if i == 0 else None)
        gap = 24
    if L["subject"]:
        D.run(ctx, box.p(before=gap, line=1.4, keep=True), L["subject"], "body", size=size,
              color=colour, bold=True)
        gap = 20
    if L["greeting"]:
        D.run(ctx, box.p(before=gap, line=1.62), L["greeting"], "body", size=size, color=colour)
        gap = 12
    for para in L["paragraphs"]:
        p = box.p(before=gap, line=1.62)
        for k, line in enumerate(para.split("\n")):
            if k:
                p.add_run().add_break()
            D.run(ctx, p, line, "body", size=size, color=colour)
        gap = 12
    if L["closing"]:
        D.run(ctx, box.p(before=22, line=1.62, keep=True), L["closing"], "body", size=size,
              color=colour)
    if L["signature"]:
        D.run(ctx, box.p(before=6 if L["closing"] else 22, line=1.2), L["signature"], "name",
              size=round(size * 1.45, 1), color=colour)


def design_letter_doc(data: dict, template_key: str):
    """(document, faces) of a Modern template's cover letter: the template's
    OWN Word design drawn with the CV content taken out - its bands, side
    column, photo, name, headline and contact exactly as on the CV - and the
    plain letter body set in its main column (user, 2026-10-20: "the
    letter's header is the CV's own header block, not an approximation")."""
    from ..labels import reset_lang, set_lang, t
    from .docx import _photo_path
    lang = lang_of(data)
    photo = _photo_path(normalize(data).get("photo_url", ""))
    full, _ = D.render(data, template_key, photo)
    jobs = normalize(data).get("experience") or [{}]
    colour = _run_colour(full, ((jobs[0].get("bullets") or [""])[0]) or "")
    # by the first BULLET: a company name can recur (a reference's "CEO,
    # Halden & Row" in modern-t20's side column)
    where = _job_column(full, ((jobs[0].get("bullets") or [""])[0])[:30]
                        or jobs[0].get("company") or "")
    doc, resolved = D.render(_emptied(data), template_key, photo)
    _drop_orphan_labels(doc)
    cell, width = _main_cell(doc, where)
    body = doc.element.body
    token = set_lang(lang)
    try:
        ctx = D.Ctx(doc=doc, r=normalize(data), lang=lang, resolved=resolved, photo=None, t=t)
        if cell is None:
            sec = doc.sections[0]
            width = int((sec.page_width - sec.left_margin - sec.right_margin) / 635)
            box = D.Box(ctx, doc, width)
        else:
            box = D.Box(ctx, cell, width)
        _write_body(ctx, box, letter_mod.letter_of(data), lang, colour)
        D.true_lines(ctx)
        D.half_leading(ctx)
    finally:
        reset_lang(token)
    if lang == "ar":
        # every paragraph: the letter's own, and the shape-carrying spacer the
        # design inserts after its RTL pass (docx_design.body_page_shapes)
        for p in body.iter(qn("w:p")):
            _put(D._ppr(p), _w("bidi"), PPR_ORDER)
    return doc, D.faces_drawn(doc, resolved)


_R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def _copy_images(el, src_part, dst_part) -> None:
    """Re-point every picture in `el` (copied from `src_part`) at the same image
    added to `dst_part` - a relationship id means nothing in another part."""
    for node in el.iter():
        for attr in (f"{_R}embed", f"{_R}id"):
            rid = node.get(attr)
            if not rid or rid not in src_part.rels:
                continue
            rel = src_part.rels[rid]
            if "image" not in rel.reltype:
                continue
            new_rid, _ = dst_part.get_or_add_image(io.BytesIO(rel.target_part.blob))
            node.set(attr, new_rid)


def append_document(cv_doc, letter_doc) -> None:
    """"CV + cover letter" in one file: the finished letter document copied
    into its OWN SECTION after the CV - a new page with the letter's page
    size, margins and header (its page shapes: a side column's fill, a
    band), so none of the CV's reach the letter and the letter's reach only
    its page. Built as a document of its own first, the letter keeps exactly
    its own fit and chosen sizes (they apply to the letter alone)."""
    import copy
    rtl = letter_doc.sections[0]._sectPr.find(qn("w:bidi")) is not None
    sec = cv_doc.add_section(WD_SECTION.NEW_PAGE)
    # the paragraph that now ends the CV's section must add nothing (see _section)
    ppr = D._ppr(cv_doc.element.body.findall(qn("w:p"))[-1])
    rpr = OxmlElement("w:rPr")
    rpr.append(_w("sz", val=2))
    rpr.append(_w("szCs", val=2))
    _put(ppr, rpr, PPR_ORDER)
    _put(ppr, _w("spacing", before=0, after=0, line=60, lineRule="auto"), PPR_ORDER)
    src = letter_doc.sections[0]
    for attr in ("page_width", "page_height", "left_margin", "right_margin", "top_margin",
                 "bottom_margin", "header_distance", "footer_distance"):
        setattr(sec, attr, getattr(src, attr))
    sec.different_first_page_header_footer = src.different_first_page_header_footer
    pairs = [(sec.header, src.header), (sec.footer, src.footer)]
    if src.different_first_page_header_footer:
        pairs.append((sec.first_page_header, src.first_page_header))
    for dst, s in pairs:
        dst.is_linked_to_previous = False
        for el in list(dst._element):
            dst._element.remove(el)
        for el in s._element:
            c = copy.deepcopy(el)
            _copy_images(c, s.part, dst.part)
            dst._element.append(c)
        if not len(dst._element):
            dst._element.append(OxmlElement("w:p"))
    body = cv_doc.element.body
    end = body[-1]                          # the document's closing sectPr
    for el in letter_doc.element.body:
        if el.tag == qn("w:sectPr"):
            continue
        c = copy.deepcopy(el)
        _copy_images(c, letter_doc.part, cv_doc.part)
        end.addprevious(c)
    if rtl:
        _section_bidi(cv_doc.sections[-1]._sectPr)


def render(data: dict, template_key: str) -> bytes:
    """The cover letter alone: a one-page, editable Word document."""
    rtl = lang_of(data) == "ar"
    doc = D.new_document(rtl)
    resolved = docx_theme.resolve(data, template_key)
    docx_theme.apply(doc, resolved)
    _build(doc, data, template_key, resolved, new_section=False)
    faces = docx_theme.faces_drawn(doc, resolved)
    buf = io.BytesIO()
    doc.save(buf)
    return docx_font_embed.embed_fonts(buf.getvalue(), faces)


def append(doc, data: dict, template_key: str) -> set:
    """Add the letter after a finished CV document; returns the faces it draws
    (the caller embeds them with the CV's)."""
    resolved = docx_theme.resolve(data, template_key)
    _build(doc, data, template_key, resolved, new_section=True)
    return docx_theme.faces_drawn(doc, resolved)
